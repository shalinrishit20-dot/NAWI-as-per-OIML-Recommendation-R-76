"""Flask integration tests for OIML R 76 observation persistence.

Uses an isolated SQLite file. Does not touch instance/nawi.db.
DATABASE_URL is set before importing app.
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest

BACKEND_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_ROOT not in sys.path:
    sys.path.insert(0, BACKEND_ROOT)

_FD, _TEST_DB = tempfile.mkstemp(prefix="nawi_r76_test_", suffix=".db")
os.close(_FD)
_TEST_ENV_KEYS = ("DATABASE_URL", "SECRET_KEY", "NAWI_DEMO_MODE", "NAWI_SESSION_COOKIE_SECURE")
_PREVIOUS_TEST_ENV = {key: os.environ.get(key) for key in _TEST_ENV_KEYS}
_HAD_TEST_ENV = {key: key in os.environ for key in _TEST_ENV_KEYS}
os.environ["DATABASE_URL"] = "sqlite:///" + _TEST_DB
os.environ["SECRET_KEY"] = "test-secret"
os.environ["NAWI_DEMO_MODE"] = "1"
os.environ["NAWI_SESSION_COOKIE_SECURE"] = "0"

try:
    from app import AuditLog, EccentricityTest, Instrument, Observation, ObservationCalculation, RepeatabilityTest, TareWeighingTest, ZeroSettingTest, _build_report_document_data, app, db  # noqa: E402
    from rules import RULESET_VERSION  # noqa: E402
    from rules.compliance import evaluate_digital_observation  # noqa: E402
finally:
    for _key in _TEST_ENV_KEYS:
        if _HAD_TEST_ENV[_key]:
            os.environ[_key] = _PREVIOUS_TEST_ENV[_key]  # type: ignore[assignment]
        else:
            os.environ.pop(_key, None)


class AppR76IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        app.config["TESTING"] = True
        cls.client = app.test_client()

    @classmethod
    def tearDownClass(cls) -> None:
        with app.app_context():
            db.session.remove()
            db.engine.dispose()
        try:
            os.unlink(_TEST_DB)
        except OSError:
            pass

    def setUp(self) -> None:
        self.client.post("/api/auth/logout")
        resp = self.client.post(
            "/api/auth/login",
            json={"username": "tester", "password": "tester123"},
        )
        self.assertEqual(resp.status_code, 200)

    def _create_instrument(self, e: float, d: float, max_capacity: float = 30000.0) -> int:
        resp = self.client.post(
            "/api/instruments",
            json={
                "manufacturer": "TestCo",
                "model": "T1",
                "serial_number": f"SN-{os.urandom(4).hex()}",
                "accuracy_class": "III",
                "max_capacity": max_capacity,
                "min_capacity": 100.0,
                "e": e,
                "d": d,
                "unit": "g",
            },
        )
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        return resp.get_json()["id"]

    def _create_report(self, instrument_id: int) -> dict:
        resp = self.client.post("/api/reports", json={"instrument_id": instrument_id})
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        return resp.get_json()

    def test_auth_and_report_ruleset_version(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        self.assertEqual(report["standard_version"], RULESET_VERSION)

    def test_direct_path_observation(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        rid = report["id"]
        expected = evaluate_digital_observation(
            "III", 5000.0, 5000.5, 5.0, 0.5, 0.2
        )
        resp = self.client.post(
            f"/api/reports/{rid}/observations",
            json={
                "applied_load": 5000.0,
                "reference_value": 5000.0,
                "indication": 5000.5,
                "zero_error": 0.2,
                "test_type": "Accuracy",
            },
        )
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        obs = resp.get_json()["observations"][0]
        self.assertEqual(obs["error"], expected.corrected_error)
        self.assertEqual(obs["mpe"], expected.mpe)
        self.assertEqual(obs["result"], expected.result)
        calc = obs["calculation"]
        self.assertIsNotNone(calc)
        self.assertEqual(calc["error_method"], "direct")
        self.assertFalse(calc["rounding_elimination_required"])
        self.assertIsNone(calc["additional_load"])
        self.assertEqual(calc["ruleset_version"], RULESET_VERSION)
        self.assertEqual(calc["canonical_load"], expected.load)
        self.assertEqual(calc["verification_scale_interval"], expected.verification_scale_interval)
        self.assertEqual(calc["actual_scale_interval"], expected.actual_scale_interval)
        with app.app_context():
            obs_rows = Observation.query.filter_by(report_id=rid).all()
            self.assertEqual(len(obs_rows), 1)
            calcs = ObservationCalculation.query.filter_by(
                observation_id=obs_rows[0].id
            ).all()
            self.assertEqual(len(calcs), 1)

    def test_successful_observation_creates_one_audit_row(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        rid = report["id"]
        with app.app_context():
            before = AuditLog.query.filter_by(report_id=rid).count()
        resp = self.client.post(
            f"/api/reports/{rid}/observations",
            json={
                "applied_load": 5000.0,
                "reference_value": 5000.0,
                "indication": 5000.5,
                "zero_error": 0.2,
            },
        )
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        result = resp.get_json()["observations"][0]["result"]
        with app.app_context():
            rows = AuditLog.query.filter_by(report_id=rid).all()
            matches = [
                a for a in rows
                if a.action == f"Added Accuracy observation: {result} [{RULESET_VERSION}]"
            ]
            self.assertEqual(len(matches), 1)
            self.assertEqual(len(rows), before + 1)

    def test_changeover_missing_delta_l_is_400_and_not_committed(self) -> None:
        with app.app_context():
            before_o = Observation.query.count()
            before_c = ObservationCalculation.query.count()
        iid = self._create_instrument(5.0, 5.0)
        report = self._create_report(iid)
        resp = self.client.post(
            f"/api/reports/{report['id']}/observations",
            json={
                "applied_load": 1000.0,
                "reference_value": 1000.0,
                "indication": 1000.0,
                "zero_error": 0.5,
            },
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Invalid observation", resp.get_json()["error"])
        with app.app_context():
            self.assertEqual(Observation.query.count(), before_o)
            self.assertEqual(ObservationCalculation.query.count(), before_c)

    def test_changeover_with_delta_l(self) -> None:
        iid = self._create_instrument(5.0, 5.0)
        report = self._create_report(iid)
        expected = evaluate_digital_observation(
            "III", 1000.0, 1000.0, 5.0, 5.0, 0.5, additional_load=1.5
        )
        resp = self.client.post(
            f"/api/reports/{report['id']}/observations",
            json={
                "applied_load": 1000.0,
                "reference_value": 1000.0,
                "indication": 1000.0,
                "zero_error": 0.5,
                "additional_load": 1.5,
            },
        )
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        calc = resp.get_json()["observations"][0]["calculation"]
        self.assertEqual(calc["error_method"], "changeover_point")
        self.assertEqual(calc["indication_before_rounding"], expected.indication_before_rounding)
        self.assertEqual(calc["error_before_rounding"], expected.error_before_rounding)
        self.assertEqual(calc["corrected_error"], expected.corrected_error)
        self.assertEqual(calc["mpe"], expected.mpe)
        self.assertEqual(calc["margin_to_limit"], expected.margin_to_limit)
        self.assertEqual(calc["additional_load"], 1.5)
        self.assertEqual(calc["error_rule_reference"], expected.error_rule_reference)
        self.assertEqual(calc["mpe_rule_reference"], expected.mpe_rule_reference)

    def test_missing_zero_error_is_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        resp = self.client.post(
            f"/api/reports/{report['id']}/observations",
            json={
                "applied_load": 5000.0,
                "reference_value": 5000.0,
                "indication": 5000.5,
            },
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("zero_error", resp.get_json().get("detail", ""))

    def test_legacy_observation_calculation_null(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        with app.app_context():
            o = Observation(
                report_id=report["id"],
                test_type="Accuracy",
                applied_load=100.0,
                reference_value=100.0,
                indication=100.0,
                error=0.0,
                mpe=1.0,
                result="PASS",
            )
            db.session.add(o)
            db.session.commit()
            oid = o.id
        resp = self.client.get(f"/api/reports/{report['id']}")
        self.assertEqual(resp.status_code, 200)
        obs = [x for x in resp.get_json()["observations"] if x["id"] == oid][0]
        self.assertIsNone(obs["calculation"])

    def test_invalid_applied_load_not_persisted(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        rid = report["id"]
        with app.app_context():
            before_o = Observation.query.filter_by(report_id=rid).count()
            before_c = ObservationCalculation.query.count()

        for bad in (True, float("nan"), float("inf"), -1.0):
            payload = {
                "applied_load": bad,
                "reference_value": 5000.0,
                "indication": 5000.5,
                "zero_error": 0.2,
            }
            resp = self.client.post(
                f"/api/reports/{rid}/observations",
                data=json.dumps(payload, allow_nan=True),
                content_type="application/json",
            )
            self.assertEqual(resp.status_code, 400, msg=repr(bad))
            self.assertEqual(resp.get_json()["error"], "Invalid observation")

        with app.app_context():
            self.assertEqual(Observation.query.filter_by(report_id=rid).count(), before_o)
            self.assertEqual(ObservationCalculation.query.count(), before_c)

    def _instrument_payload(self, **overrides):
        body = {
            "manufacturer": "TestCo",
            "model": "T1",
            "serial_number": f"SN-{os.urandom(4).hex()}",
            "accuracy_class": "III",
            "max_capacity": "30",
            "min_capacity": "0.10",
            "e": "0.01",
            "d": "0.005",
            "unit": "kg",
        }
        body.update(overrides)
        return body

    def test_valid_class_iii_string_numerics(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
        payload = self._instrument_payload()
        resp = self.client.post("/api/instruments", json=payload)
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        body = resp.get_json()
        c = body["classification"]
        self.assertTrue(c["valid"])
        self.assertEqual(c["number_of_verification_intervals"], 3000.0)
        self.assertEqual(c["minimum_capacity_multiplier"], 20.0)
        self.assertEqual(c["minimum_capacity_required"], 0.10)
        self.assertEqual(c["minimum_capacity_in_d"], 20.0)
        self.assertEqual(c["rule_reference"], "OIML R 76-1:2006 §3.2 Table 3; §3.4.3")
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before + 1)

    def test_min_capacity_classification_failure_422(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
            before_audit = AuditLog.query.filter(
                AuditLog.action.like("Registered instrument%")
            ).count()
        resp = self.client.post(
            "/api/instruments", json=self._instrument_payload(min_capacity="0.099")
        )
        self.assertEqual(resp.status_code, 422)
        body = resp.get_json()
        self.assertIn("OIML R76 classification", body["error"])
        c = body["classification"]
        self.assertFalse(c["valid"])
        self.assertTrue(c["errors"])
        self.assertEqual(c["minimum_capacity_required"], 0.10)
        self.assertTrue(any("19.8" in e for e in c["errors"]))
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before)
            self.assertEqual(
                AuditLog.query.filter(
                    AuditLog.action.like("Registered instrument%")
                ).count(),
                before_audit,
            )

    def test_n_above_class_iii_limit_422(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
        resp = self.client.post(
            "/api/instruments",
            json=self._instrument_payload(max_capacity="200", min_capacity="0.10"),
        )
        self.assertEqual(resp.status_code, 422)
        c = resp.get_json()["classification"]
        self.assertFalse(c["valid"])
        self.assertTrue(any("10000" in e or "10,000" in e or "n =" in e for e in c["errors"]))
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before)

    def test_malformed_instrument_numeric_400(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
            before_audit = AuditLog.query.filter(
                AuditLog.action.like("Registered instrument%")
            ).count()
        cases = [
            {"max_capacity": True},
            {"e": float("nan")},
            {"d": float("inf")},
            {"min_capacity": float("-inf")},
            {"max_capacity": "not-a-number"},
        ]
        for extra in cases:
            payload = self._instrument_payload(**extra)
            resp = self.client.post(
                "/api/instruments",
                data=json.dumps(payload, allow_nan=True),
                content_type="application/json",
            )
            self.assertEqual(resp.status_code, 400, msg=repr(extra))
            self.assertEqual(resp.get_json()["error"], "Invalid instrument data")
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before)
            self.assertEqual(
                AuditLog.query.filter(
                    AuditLog.action.like("Registered instrument%")
                ).count(),
                before_audit,
            )

    def test_identity_and_empty_class_are_400(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
            before_audit = AuditLog.query.filter(
                AuditLog.action.like("Registered instrument%")
            ).count()
        r1 = self.client.post(
            "/api/instruments", json=self._instrument_payload(manufacturer=123)
        )
        self.assertEqual(r1.status_code, 400)
        r2 = self.client.post(
            "/api/instruments", json=self._instrument_payload(model="   ")
        )
        self.assertEqual(r2.status_code, 400)
        r3 = self.client.post(
            "/api/instruments", json=self._instrument_payload(accuracy_class="")
        )
        self.assertEqual(r3.status_code, 400)
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before)
            self.assertEqual(
                AuditLog.query.filter(
                    AuditLog.action.like("Registered instrument%")
                ).count(),
                before_audit,
            )

    def test_unsupported_class_and_unit_are_400(self) -> None:
        with app.app_context():
            before = Instrument.query.count()
        r1 = self.client.post(
            "/api/instruments", json=self._instrument_payload(accuracy_class="IV")
        )
        self.assertEqual(r1.status_code, 400)
        r2 = self.client.post(
            "/api/instruments", json=self._instrument_payload(unit="mg")
        )
        self.assertEqual(r2.status_code, 400)
        with app.app_context():
            self.assertEqual(Instrument.query.count(), before)

    def test_successful_instrument_creates_one_audit_row(self) -> None:
        payload = self._instrument_payload()
        resp = self.client.post("/api/instruments", json=payload)
        self.assertEqual(resp.status_code, 201, resp.get_data(as_text=True))
        iid = resp.get_json()["id"]
        expected = f"Registered instrument {iid} [{RULESET_VERSION} classification valid]"
        with app.app_context():
            matches = AuditLog.query.filter_by(action=expected).all()
            self.assertEqual(len(matches), 1)



    def _put_repeatability(self, rid, role, load, readings):
        return self.client.put(
            f"/api/reports/{rid}/repeatability/{role}",
            json={"applied_load": load, "readings": readings},
        )

    def _put_eccentricity(self, rid, positions, tare=0):
        body = {"position_errors": positions}
        if tare is not None:
            body["maximum_additive_tare_effect"] = tare
        return self.client.put(f"/api/reports/{rid}/eccentricity", json=body)

    def test_repeatability_half_max_pass_and_audit(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before = AuditLog.query.filter_by(report_id=rid).count()
        readings = [15000.0] * 10
        resp = self._put_repeatability(rid, "half_max", 15000.0, readings)
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["series_role"], "half_max")
        self.assertEqual(body["status"], "PASS")
        self.assertEqual(body["required_weighings"], 10)
        self.assertEqual(body["weighing_count"], 10)
        with app.app_context():
            rows = RepeatabilityTest.query.filter_by(report_id=rid).all()
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].series_role, "half_max")
            audits = [
                a for a in AuditLog.query.filter_by(report_id=rid).all()
                if a.action == f"Evaluated repeatability half_max [{RULESET_VERSION}: PASS]"
            ]
            self.assertEqual(len(audits), 1)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before + 1)

    def test_repeatability_near_max_independent(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.assertEqual(self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10).status_code, 200)
        resp = self._put_repeatability(rid, "near_max", 30000.0, [30000.0] * 10)
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        self.assertEqual(resp.get_json()["series_role"], "near_max")
        with app.app_context():
            roles = {x.series_role for x in RepeatabilityTest.query.filter_by(report_id=rid).all()}
            self.assertEqual(roles, {"half_max", "near_max"})

    def test_repeatability_resubmit_same_role(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10)
        resp = self._put_repeatability(rid, "half_max", 14000.0, [14000.0] * 10)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["applied_load"], 14000.0)
        with app.app_context():
            self.assertEqual(RepeatabilityTest.query.filter_by(report_id=rid, series_role="half_max").count(), 1)

    def test_repeatability_incomplete_stored(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 3)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["status"], "INCOMPLETE")
        self.assertFalse(resp.get_json()["complete"])

    def test_repeatability_malformed_400_no_write(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = RepeatabilityTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        for payload in (
            {"applied_load": 15000.0, "readings": [True]},
            {"applied_load": 15000.0, "readings": [float("nan")]},
            {"applied_load": True, "readings": [15000.0]},
        ):
            resp = self.client.put(
                f"/api/reports/{rid}/repeatability/half_max",
                data=json.dumps(payload, allow_nan=True),
                content_type="application/json",
            )
            self.assertEqual(resp.status_code, 400, msg=repr(payload))
        with app.app_context():
            self.assertEqual(RepeatabilityTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_eccentricity_four_positions_pass(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["status"], "PASS")
        self.assertEqual(body["required_test_load"], 10000.0)
        self.assertEqual(body["ruleset_version"], RULESET_VERSION)
        with app.app_context():
            self.assertEqual(EccentricityTest.query.filter_by(report_id=rid).count(), 1)

    def test_eccentricity_partial_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["status"], "INCOMPLETE")

    def test_eccentricity_unknown_position_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = EccentricityTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_eccentricity(rid, {"Q1": 0.0, "Q5": 0.0})
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(EccentricityTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_eccentricity_resubmit_updates(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_eccentricity(rid, {"Q1": 0.0})
        resp = self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["status"], "PASS")
        with app.app_context():
            self.assertEqual(EccentricityTest.query.filter_by(report_id=rid).count(), 1)

    def test_programme_empty_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        prog = report["test_programme"]
        self.assertEqual(prog["overall_status"], "INCOMPLETE")
        self.assertFalse(prog["full_standard_coverage"])
        self.assertIsNone(report["repeatability"]["half_max"])
        self.assertIsNone(report["repeatability"]["near_max"])
        self.assertIsNone(report["eccentricity"])
        self.assertIsNone(report["zero_setting"])
        self.assertIsNone(report.get("tare_weighing"))

    def test_programme_observation_only_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.client.post(
            f"/api/reports/{rid}/observations",
            json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2},
        )
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertEqual(prog["observation_status"], "PASS")
        self.assertEqual(prog["overall_status"], "INCOMPLETE")
        self.assertFalse(prog["full_standard_coverage"])

    def _fill_core_pass(self, rid):
        self.client.post(
            f"/api/reports/{rid}/observations",
            json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2},
        )
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10)
        self._put_repeatability(rid, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])

    def test_programme_all_pass_not_full_coverage(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        body = self.client.get(f"/api/reports/{rid}").get_json()
        prog = body["test_programme"]
        self.assertEqual(prog["overall_status"], "PASS")
        self.assertFalse(prog["full_standard_coverage"])
        self.assertTrue(prog["rule_references"])
        self.assertEqual(prog["legacy_observation_count"], 0)

    def test_programme_repeatability_fail(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.client.post(
            f"/api/reports/{rid}/observations",
            json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2},
        )
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 9 + [15020.0])
        self._put_repeatability(rid, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertEqual(prog["repeatability_status"], "FAIL")
        self.assertEqual(prog["overall_status"], "FAIL")

    def test_programme_eccentricity_fail(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        self._put_eccentricity(rid, {"Q1": 20.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertEqual(prog["eccentricity_status"], "FAIL")
        self.assertEqual(prog["overall_status"], "FAIL")

    def test_legacy_observation_blocks_observation_pass(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.client.post(
            f"/api/reports/{rid}/observations",
            json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2},
        )
        with app.app_context():
            db.session.add(Observation(
                report_id=rid, test_type="Accuracy", applied_load=100.0,
                reference_value=100.0, indication=100.0, error=0.0, mpe=1.0, result="PASS",
            ))
            db.session.commit()
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertGreater(prog["legacy_observation_count"], 0)
        self.assertNotEqual(prog["observation_status"], "PASS")
        self.assertEqual(prog["observation_status"], "INCOMPLETE")

    def test_get_does_not_rerun_standards_engines(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("standards engine should not rerun on GET")
        orig_r = app_module.evaluate_type_approval_repeatability_series
        orig_e = app_module.evaluate_quarter_segment_eccentricity
        orig_o = app_module.evaluate_digital_observation
        orig_z = app_module.evaluate_zero_setting_accuracy
        orig_t = app_module.evaluate_tare_weighing_series
        app_module.evaluate_type_approval_repeatability_series = boom
        app_module.evaluate_quarter_segment_eccentricity = boom
        app_module.evaluate_digital_observation = boom
        app_module.evaluate_zero_setting_accuracy = boom
        app_module.evaluate_tare_weighing_series = boom
        try:
            resp = self.client.get(f"/api/reports/{rid}")
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(resp.get_json()["test_programme"]["overall_status"], "PASS")
            self.assertFalse(resp.get_json()["test_programme"]["full_standard_coverage"])
            self.assertEqual(resp.get_json()["test_programme"]["zero_setting_status"], "PASS")
            self.assertEqual(resp.get_json()["test_programme"]["tare_weighing_status"], "PASS")
        finally:
            app_module.evaluate_type_approval_repeatability_series = orig_r
            app_module.evaluate_quarter_segment_eccentricity = orig_e
            app_module.evaluate_digital_observation = orig_o
            app_module.evaluate_zero_setting_accuracy = orig_z
            app_module.evaluate_tare_weighing_series = orig_t




    def _docx_text(self, path):
        from docx import Document
        d = Document(path)
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    def _cleanup_generated_reports(self, report_number: str) -> None:
        reports_dir = os.path.join(BACKEND_ROOT, "reports")
        for ext in (".pdf", ".docx"):
            path = os.path.join(reports_dir, report_number + ext)
            try:
                os.unlink(path)
            except OSError:
                pass

    def test_report_documents_empty_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        pdf = self.client.get(f"/api/reports/{rid}/pdf")
        docx = self.client.get(f"/api/reports/{rid}/word")
        try:
            self.assertEqual(pdf.status_code, 200)
            self.assertTrue((pdf.mimetype or "").startswith("application/pdf") or pdf.headers.get("Content-Type", "").startswith("application/pdf"))
            self.assertGreater(len(pdf.data), 100)
            self.assertEqual(docx.status_code, 200)
            ctype = docx.headers.get("Content-Type", "")
            self.assertIn("officedocument.wordprocessingml.document", ctype)
            self.assertGreater(len(docx.data), 100)
            with app.app_context():
                r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
                data = _build_report_document_data(r)
            self.assertFalse(data["programme"]["full_standard_coverage"])
            self.assertIn("not a complete", data["scope_warning"].lower())
            self.assertEqual(data["instrument"]["manufacturer"], "TestCo")
            self.assertEqual(data["instrument"]["n"], "Not recorded")
            self.assertIsNone(data["repeatability"]["half_max"])
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Implemented Core Test Report", text)
            self.assertIn("full_standard_coverage", text)
            self.assertIn("Not recorded — test incomplete", text)
            self.assertNotIn("Type Approved", text)
            self.assertNotIn("Certified", text)
        finally:
            pdf.close()
            docx.close()
            self._cleanup_generated_reports(number)

    def test_report_documents_snapshots_and_legacy(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        self.client.post(
            f"/api/reports/{rid}/observations",
            json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2},
        )
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10)
        self._put_eccentricity(rid, {"Q1": 0.1, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        with app.app_context():
            db.session.add(Observation(
                report_id=rid, test_type="Accuracy", applied_load=100.0,
                reference_value=100.0, indication=100.0, error=0.0, mpe=1.0, result="PASS",
            ))
            db.session.commit()
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        snap = next(o["calculation"] for o in data["observations"] if not o["legacy"])
        self.assertIn("absolute_corrected_error", snap)
        self.assertIn("margin_to_limit", snap)
        self.assertIn("error_method", snap)
        self.assertIn("ruleset_version", snap)
        self.assertIn("error_rule_reference", snap)
        self.assertIn("mpe_rule_reference", snap)
        self.assertEqual(data["instrument"]["n"], "Not recorded")
        self.assertTrue(any(o["legacy"] for o in data["observations"]))
        self.assertEqual(data["repeatability"]["half_max"]["status"], "PASS")
        self.assertIsNone(data["repeatability"]["near_max"])
        self.assertEqual(data["eccentricity"]["position_errors"]["Q1"], 0.1)
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("standards engine must not rerun")
        orig = (
            app_module.evaluate_digital_observation,
            app_module.evaluate_type_approval_repeatability_series,
            app_module.evaluate_quarter_segment_eccentricity,
        )
        app_module.evaluate_digital_observation = boom
        app_module.evaluate_type_approval_repeatability_series = boom
        app_module.evaluate_quarter_segment_eccentricity = boom
        try:
            pdf = self.client.get(f"/api/reports/{rid}/pdf")
            docx = self.client.get(f"/api/reports/{rid}/word")
        finally:
            app_module.evaluate_digital_observation, app_module.evaluate_type_approval_repeatability_series, app_module.evaluate_quarter_segment_eccentricity = orig
        try:
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(docx.status_code, 200)
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Legacy / no OIML calculation snapshot", text)
            self.assertIn("15000", text)
            self.assertIn("Q1 Ec", text)
            self.assertIn("R76-1-2006-v1", text)
            self.assertNotIn("Type Approved", text)
            self.assertNotIn("Certified", text)
        finally:
            pdf.close()
            docx.close()
            self._cleanup_generated_reports(number)




    def _login(self, username, password):
        self.client.post("/api/auth/logout")
        resp = self.client.post("/api/auth/login", json={"username": username, "password": password})
        self.assertEqual(resp.status_code, 200)

    def test_incomplete_cannot_complete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before = AuditLog.query.filter_by(report_id=rid).count()
        resp = self.client.post(f"/api/reports/{rid}/complete")
        self.assertEqual(resp.status_code, 409)
        body = resp.get_json()
        self.assertEqual(body["overall_status"], "INCOMPLETE")
        self.assertTrue(body["incomplete_modules"])
        self.assertIn("scope_note", body)
        with app.app_context():
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before)
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            self.assertEqual(r.status, "In Progress")

    def test_complete_pass_and_fail_programmes(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        resp = self.client.post(f"/api/reports/{rid}/complete")
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["status"], "Review")
        self.assertEqual(body["final_result"], "PASS")
        self.assertEqual(body["test_programme"]["overall_status"], "PASS")
        with app.app_context():
            audits = [a for a in AuditLog.query.filter_by(report_id=rid).all() if a.action == "Submitted for review: PASS"]
            self.assertEqual(len(audits), 1)
        again = self.client.post(f"/api/reports/{rid}/complete")
        self.assertEqual(again.status_code, 409)

        rid2 = self._create_report(iid)["id"]
        self.client.post(f"/api/reports/{rid2}/observations", json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2})
        self._put_repeatability(rid2, "half_max", 15000.0, [15000.0] * 9 + [15020.0])
        self._put_repeatability(rid2, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(rid2, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        resp = self.client.post(f"/api/reports/{rid2}/complete")
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        self.assertEqual(resp.get_json()["final_result"], "FAIL")
        self.assertEqual(resp.get_json()["status"], "Review")

    def test_approve_preconditions_and_rbac(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        draft = self._create_report(iid)["id"]
        self._login("reviewer", "reviewer123")
        resp = self.client.post(f"/api/reports/{draft}/approve")
        self.assertEqual(resp.status_code, 409)
        self._login("tester", "tester123")
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        self.client.post(f"/api/reports/{rid}/complete")
        denied = self.client.post(f"/api/reports/{rid}/approve")
        self.assertEqual(denied.status_code, 403)
        self._login("reviewer", "reviewer123")
        ok = self.client.post(f"/api/reports/{rid}/approve")
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(ok.get_json()["status"], "Completed")
        self.assertEqual(ok.get_json()["final_result"], "PASS")
        with app.app_context():
            audits = [a for a in AuditLog.query.filter_by(report_id=rid).all() if a.action == "Report approved"]
            self.assertEqual(len(audits), 1)
        again = self.client.post(f"/api/reports/{rid}/approve")
        self.assertEqual(again.status_code, 409)
        complete_after = self.client.post(f"/api/reports/{rid}/complete")
        self.assertEqual(complete_after.status_code, 403)

        self._login("tester", "tester123")
        complete_after_authorized = self.client.post(
            f"/api/reports/{rid}/complete"
        )
        self.assertEqual(complete_after_authorized.status_code, 409)

        self._login("tester", "tester123")
        ridf = self._create_report(iid)["id"]
        self.client.post(f"/api/reports/{ridf}/observations", json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2})
        self._put_repeatability(ridf, "half_max", 15000.0, [15000.0] * 9 + [15020.0])
        self._put_repeatability(ridf, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(ridf, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self.client.post(f"/api/reports/{ridf}/complete")
        self._login("admin", "admin123")
        okf = self.client.post(f"/api/reports/{ridf}/approve")
        self.assertEqual(okf.status_code, 200)
        self.assertEqual(okf.get_json()["final_result"], "FAIL")
        self.assertEqual(okf.get_json()["status"], "Completed")

    def test_writes_locked_after_completion(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        env_ok = self.client.post(
            f"/api/reports/{rid}/environment",
            json={"temperature": 20.0, "humidity": 48.0, "pressure": 1013, "laboratory": "NAWI Lab", "location": "Floor"},
        )
        self.assertEqual(env_ok.status_code, 200)
        self.client.post(f"/api/reports/{rid}/complete")
        with app.app_context():
            before_o = Observation.query.filter_by(report_id=rid).count()
            before_r = RepeatabilityTest.query.filter_by(report_id=rid, series_role="half_max").first()
            half_load = before_r.applied_load
            before_e = EccentricityTest.query.filter_by(report_id=rid).first().required_test_load
        obs = self.client.post(f"/api/reports/{rid}/observations", json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2})
        self.assertEqual(obs.status_code, 409)
        rep = self._put_repeatability(rid, "half_max", 14000.0, [14000.0] * 10)
        self.assertEqual(rep.status_code, 409)
        ecc = self._put_eccentricity(rid, {"Q1": 1.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self.assertEqual(ecc.status_code, 409)
        with app.app_context():
            self.assertEqual(Observation.query.filter_by(report_id=rid).count(), before_o)
            self.assertEqual(RepeatabilityTest.query.filter_by(report_id=rid, series_role="half_max").first().applied_load, half_load)
            self.assertEqual(EccentricityTest.query.filter_by(report_id=rid).first().required_test_load, before_e)
        got = self.client.get(f"/api/reports/{rid}")
        self.assertEqual(got.status_code, 200)
        self.assertEqual(got.get_json()["test_programme"]["overall_status"], "PASS")
        pdf = self.client.get(f"/api/reports/{rid}/pdf")
        docx = self.client.get(f"/api/reports/{rid}/word")
        try:
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(docx.status_code, 200)
        finally:
            pdf.close(); docx.close()
            self._cleanup_generated_reports(got.get_json()["report_number"])

    def test_complete_uses_snapshots_not_engines(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("standards engine must not rerun")
        orig = (
            app_module.evaluate_digital_observation,
            app_module.evaluate_type_approval_repeatability_series,
            app_module.evaluate_quarter_segment_eccentricity,
            app_module.evaluate_zero_setting_accuracy,
            app_module.evaluate_tare_weighing_series,
        )
        app_module.evaluate_digital_observation = boom
        app_module.evaluate_type_approval_repeatability_series = boom
        app_module.evaluate_quarter_segment_eccentricity = boom
        app_module.evaluate_zero_setting_accuracy = boom
        app_module.evaluate_tare_weighing_series = boom
        try:
            resp = self.client.post(f"/api/reports/{rid}/complete")
            self.assertEqual(resp.status_code, 200)
            self._login("reviewer", "reviewer123")
            appr = self.client.post(f"/api/reports/{rid}/approve")
            self.assertEqual(appr.status_code, 200)
        finally:
            (app_module.evaluate_digital_observation, app_module.evaluate_type_approval_repeatability_series, app_module.evaluate_quarter_segment_eccentricity, app_module.evaluate_zero_setting_accuracy, app_module.evaluate_tare_weighing_series) = orig

    def _put_zero_setting(self, rid, body):
        return self.client.put(f"/api/reports/{rid}/zero-setting", json=body)

    def test_zero_setting_absent_is_null(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        self.assertIsNone(report["zero_setting"])
        got = self.client.get(f"/api/reports/{report['id']}")
        self.assertIsNone(got.get_json()["zero_setting"])

    def test_zero_setting_direct_pass(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["status"], "PASS")
        self.assertEqual(body["error_method"], "direct")
        self.assertEqual(body["zero_error"], 0.0)
        self.assertEqual(body["verification_scale_interval"], 5.0)
        self.assertEqual(body["actual_scale_interval"], 0.5)
        self.assertEqual(body["unit"], "g")
        self.assertIsNone(body["additional_load"])
        got = self.client.get(f"/api/reports/{rid}").get_json()["zero_setting"]
        self.assertEqual(got["status"], "PASS")
        self.assertEqual(got["zero_error"], 0.0)
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.filter_by(report_id=rid).count(), 1)
            audits = [a for a in AuditLog.query.filter_by(report_id=rid).all() if a.action == "Saved zero-setting accuracy: PASS"]
            self.assertEqual(len(audits), 1)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before + 1)

    def test_zero_setting_direct_fail(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_zero_setting(rid, {"setting_mode": "semi_automatic", "indication": 2.0})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["status"], "FAIL")
        self.assertGreater(body["absolute_zero_error"], body["permitted_limit"])
        with app.app_context():
            audits = [a for a in AuditLog.query.filter_by(report_id=rid).all() if a.action == "Saved zero-setting accuracy: FAIL"]
            self.assertEqual(len(audits), 1)

    def test_zero_setting_changeover_persists(self) -> None:
        iid = self._create_instrument(5.0, 5.0)
        rid = self._create_report(iid)["id"]
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0, "additional_load": 2.5})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["error_method"], "changeover_point")
        self.assertEqual(body["additional_load"], 2.5)
        self.assertEqual(body["indication_before_rounding"], 0.0)
        self.assertEqual(body["zero_error"], 0.0)
        self.assertEqual(body["status"], "PASS")
        self.assertIn("A.4.4.3", body["rule_reference"])
        self.assertEqual(body["ruleset_version"], RULESET_VERSION)

    def test_zero_setting_resubmit_same_row(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.1})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["indication"], 0.1)
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.filter_by(report_id=rid).count(), 1)

    def test_zero_setting_missing_delta_l_400(self) -> None:
        iid = self._create_instrument(5.0, 5.0)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = ZeroSettingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.get_json()["error"], "Invalid zero-setting data")
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_zero_setting_delta_l_on_direct_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = ZeroSettingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0, "additional_load": 0.1})
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_zero_setting_unsupported_mode_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = ZeroSettingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_zero_setting(rid, {"setting_mode": "automatic", "indication": 0.0})
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_zero_setting_ignores_client_e_d_unit(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_zero_setting(rid, {
            "setting_mode": "non_automatic",
            "indication": 0.0,
            "e": 99.0,
            "d": 99.0,
            "unit": "t",
            "ruleset_version": "forged",
        })
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["verification_scale_interval"], 5.0)
        self.assertEqual(body["actual_scale_interval"], 0.5)
        self.assertEqual(body["unit"], "g")
        self.assertEqual(body["ruleset_version"], RULESET_VERSION)

    def test_zero_setting_locked_in_review_and_completed(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        self.client.post(f"/api/reports/{rid}/complete")
        with app.app_context():
            before = ZeroSettingTest.query.filter_by(report_id=rid).count()
        resp = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        self.assertEqual(resp.status_code, 409)
        with app.app_context():
            self.assertEqual(ZeroSettingTest.query.filter_by(report_id=rid).count(), before)
        self._login("reviewer", "reviewer123")
        self.client.post(f"/api/reports/{rid}/approve")
        self._login("tester", "tester123")
        resp2 = self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        self.assertEqual(resp2.status_code, 409)

    def test_zero_setting_get_does_not_rerun_engine(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("zero-setting engine must not rerun on GET")
        orig = app_module.evaluate_zero_setting_accuracy
        app_module.evaluate_zero_setting_accuracy = boom
        try:
            resp = self.client.get(f"/api/reports/{rid}")
            self.assertEqual(resp.status_code, 200)
            zs = resp.get_json()["zero_setting"]
            self.assertEqual(zs["status"], "PASS")
            self.assertEqual(zs["zero_error"], 0.0)
            prog = resp.get_json()["test_programme"]
            self.assertEqual(prog["zero_setting_status"], "PASS")
        finally:
            app_module.evaluate_zero_setting_accuracy = orig


    def test_report_documents_include_zero_setting_snapshot(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        with app.app_context():
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        self.assertIsNone(data["zero_setting"])
        self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        with app.app_context():
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        self.assertIsNotNone(data["zero_setting"])
        self.assertEqual(data["zero_setting"]["status"], "PASS")
        self.assertEqual(data["zero_setting"]["zero_error"], 0.0)
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("zero-setting engine must not rerun for documents")
        orig = app_module.evaluate_zero_setting_accuracy
        app_module.evaluate_zero_setting_accuracy = boom
        try:
            pdf = self.client.get(f"/api/reports/{rid}/pdf")
            docx = self.client.get(f"/api/reports/{rid}/word")
        finally:
            app_module.evaluate_zero_setting_accuracy = orig
        try:
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(docx.status_code, 200)
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Zero-setting accuracy", text)
            self.assertIn("non_automatic", text)
            self.assertNotIn("Type Approved", text)
        finally:
            pdf.close(); docx.close()
            self._cleanup_generated_reports(number)

    def test_report_documents_absent_zero_setting_not_recorded(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        docx = self.client.get(f"/api/reports/{rid}/word")
        try:
            self.assertEqual(docx.status_code, 200)
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Zero-setting accuracy", text)
            self.assertIn("Not recorded — test not performed", text)
        finally:
            docx.close()
            self._cleanup_generated_reports(number)



    def _put_tare(self, rid, tare, observations):
        return self.client.put(f"/api/reports/{rid}/tare-weighing", json={"tare_value": tare, "observations": observations})

    def _tare_obs(self, net, indication, zero_error=0.2, additional_load=None):
        body = {"net_load": net, "reference_value": net, "indication": indication, "zero_error": zero_error}
        if additional_load is not None:
            body["additional_load"] = additional_load
        return body

    def test_tare_weighing_absent_is_null(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        report = self._create_report(iid)
        self.assertIsNone(report["tare_weighing"])

    def test_tare_weighing_five_point_pass(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before = AuditLog.query.filter_by(report_id=rid).count()
        obs = [self._tare_obs(5000.0, 5000.5) for _ in range(5)]
        resp = self._put_tare(rid, 0.0, obs)
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["status"], "PASS")
        self.assertEqual(body["observation_count"], 5)
        self.assertEqual(body["tare_value"], 0.0)
        self.assertEqual(body["observations"][0]["gross_load"], 5000.0)
        self.assertIn("corrected_error", body["observations"][0])
        self.assertIn("mpe", body["observations"][0])
        self.assertIn("margin_to_limit", body["observations"][0])
        self.assertEqual(body["observations"][0]["result"], "PASS")
        got = self.client.get(f"/api/reports/{rid}").get_json()["tare_weighing"]
        self.assertEqual(got["status"], "PASS")
        self.assertEqual(self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]["tare_weighing_status"], "PASS")
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.filter_by(report_id=rid).count(), 1)
            audits = [a for a in AuditLog.query.filter_by(report_id=rid).all() if a.action == "Saved tare weighing: PASS"]
            self.assertEqual(len(audits), 1)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before + 1)

    def test_tare_weighing_fail_indices(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        obs = [self._tare_obs(5000.0, 5000.5) for _ in range(4)]
        obs.append(self._tare_obs(5000.0, 5200.0, 0.0))
        resp = self._put_tare(rid, 0.0, obs)
        self.assertEqual(resp.status_code, 200)
        body = resp.get_json()
        self.assertEqual(body["status"], "FAIL")
        self.assertEqual(body["failed_observation_indices"], [4])

    def test_tare_weighing_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(3)])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["status"], "INCOMPLETE")
        self.assertFalse(resp.get_json()["complete"])

    def test_tare_weighing_positive_tare_gross(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self._put_tare(rid, 1000.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["tare_value"], 1000.0)
        self.assertEqual(resp.get_json()["observations"][0]["gross_load"], 6000.0)

    def test_tare_weighing_changeover_snapshot(self) -> None:
        iid = self._create_instrument(5.0, 5.0)
        rid = self._create_report(iid)["id"]
        obs = [self._tare_obs(1000.0, 1000.0, 0.5, additional_load=1.5) for _ in range(5)]
        resp = self._put_tare(rid, 0.0, obs)
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        snap = resp.get_json()["observations"][0]
        self.assertEqual(snap["error_method"], "changeover_point")
        self.assertEqual(snap["additional_load"], 1.5)
        self.assertIn("indication_before_rounding", snap)

    def test_tare_weighing_resubmit_same_row(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        resp = self._put_tare(rid, 200.0, [self._tare_obs(4000.0, 4000.0, 0.0) for _ in range(5)])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["tare_value"], 200.0)
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.filter_by(report_id=rid).count(), 1)

    def test_tare_weighing_gross_over_max_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = TareWeighingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_tare(rid, 10000.0, [self._tare_obs(25000.0, 25000.0, 0.0)])
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_tare_weighing_negative_tare_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = TareWeighingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self._put_tare(rid, -1.0, [self._tare_obs(100.0, 100.0, 0.0)])
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_tare_weighing_malformed_400(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        with app.app_context():
            before_t = TareWeighingTest.query.count()
            before_a = AuditLog.query.filter_by(report_id=rid).count()
        resp = self.client.put(
            f"/api/reports/{rid}/tare-weighing",
            data=json.dumps({"tare_value": 0.0, "observations": [{"net_load": float("nan"), "indication": 1, "zero_error": 0}]}, allow_nan=True),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 400)
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.count(), before_t)
            self.assertEqual(AuditLog.query.filter_by(report_id=rid).count(), before_a)

    def test_tare_weighing_ignores_client_instrument_override(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        resp = self.client.put(f"/api/reports/{rid}/tare-weighing", json={
            "tare_value": 0.0,
            "e": 99, "d": 99, "unit": "t", "accuracy_class": "I", "max_capacity": 1,
            "observations": [self._tare_obs(5000.0, 5000.5) for _ in range(5)],
        })
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        body = resp.get_json()
        self.assertEqual(body["verification_scale_interval"], 5.0)
        self.assertEqual(body["actual_scale_interval"], 0.5)
        self.assertEqual(body["unit"], "g")
        self.assertEqual(body["accuracy_class"], "III")
        self.assertEqual(body["maximum_capacity"], 30000.0)

    def test_tare_weighing_locked_review_and_completed(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._fill_core_pass(rid)
        self.client.post(f"/api/reports/{rid}/complete")
        with app.app_context():
            before = TareWeighingTest.query.filter_by(report_id=rid).count()
        resp = self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        self.assertEqual(resp.status_code, 409)
        with app.app_context():
            self.assertEqual(TareWeighingTest.query.filter_by(report_id=rid).count(), before)
        self._login("reviewer", "reviewer123")
        self.client.post(f"/api/reports/{rid}/approve")
        self._login("tester", "tester123")
        resp2 = self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        self.assertEqual(resp2.status_code, 409)

    def test_tare_weighing_get_does_not_rerun_engine(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self._put_tare(rid, 0.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("tare/compliance engine must not rerun on GET")
        orig_t = app_module.evaluate_tare_weighing_series
        orig_o = app_module.evaluate_digital_observation
        app_module.evaluate_tare_weighing_series = boom
        app_module.evaluate_digital_observation = boom
        try:
            resp = self.client.get(f"/api/reports/{rid}")
            self.assertEqual(resp.status_code, 200)
            tw = resp.get_json()["tare_weighing"]
            self.assertEqual(tw["status"], "PASS")
            self.assertEqual(tw["observation_count"], 5)
        finally:
            app_module.evaluate_tare_weighing_series = orig_t
            app_module.evaluate_digital_observation = orig_o



    def test_report_documents_include_tare_weighing(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        with app.app_context():
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        self.assertIsNone(data.get("tare_weighing"))
        self._put_tare(rid, 1000.0, [self._tare_obs(5000.0, 5000.5) for _ in range(5)])
        with app.app_context():
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        self.assertEqual(data["tare_weighing"]["status"], "PASS")
        self.assertEqual(data["tare_weighing"]["tare_value"], 1000.0)
        self.assertEqual(data["tare_weighing"]["observations"][0]["gross_load"], 6000.0)
        import app as app_module
        def boom(*a, **k):
            raise AssertionError("tare/compliance engine must not rerun for documents")
        orig_t = app_module.evaluate_tare_weighing_series
        orig_o = app_module.evaluate_digital_observation
        app_module.evaluate_tare_weighing_series = boom
        app_module.evaluate_digital_observation = boom
        try:
            pdf = self.client.get(f"/api/reports/{rid}/pdf")
            docx = self.client.get(f"/api/reports/{rid}/word")
        finally:
            app_module.evaluate_tare_weighing_series = orig_t
            app_module.evaluate_digital_observation = orig_o
        try:
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(docx.status_code, 200)
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Tare weighing", text)
            self.assertIn("1000", text)
            self.assertIn("6000", text)
            self.assertNotIn("Type Approved", text)
        finally:
            pdf.close(); docx.close()
            self._cleanup_generated_reports(number)

    def test_report_documents_absent_tare_not_recorded(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        number = created["report_number"]
        docx = self.client.get(f"/api/reports/{rid}/word")
        try:
            self.assertEqual(docx.status_code, 200)
            fd, path = tempfile.mkstemp(suffix=".docx")
            os.close(fd)
            try:
                with open(path, "wb") as f:
                    f.write(docx.data)
                text = self._docx_text(path)
            finally:
                os.unlink(path)
            self.assertIn("Tare weighing", text)
            self.assertIn("Not recorded — test not performed", text)
        finally:
            docx.close()
            self._cleanup_generated_reports(number)

    def test_report_documents_tare_fail_status(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        created = self._create_report(iid)
        rid = created["id"]
        obs = [self._tare_obs(5000.0, 5000.5) for _ in range(4)]
        obs.append(self._tare_obs(5000.0, 5200.0, 0.0))
        self._put_tare(rid, 0.0, obs)
        with app.app_context():
            r = db.session.get(__import__("app", fromlist=["Report"]).Report, rid)
            data = _build_report_document_data(r)
        self.assertEqual(data["tare_weighing"]["status"], "FAIL")
        self.assertEqual(data["tare_weighing"]["failed_observation_indices"], [4])


    def test_programme_three_modules_without_zero_tare_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.client.post(f"/api/reports/{rid}/observations", json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2})
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10)
        self._put_repeatability(rid, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertEqual(prog["overall_status"], "INCOMPLETE")
        self.assertIn("zero_setting", prog["incomplete_modules"])
        self.assertIn("tare_weighing", prog["incomplete_modules"])
        self.assertFalse(prog["full_standard_coverage"])
        blocked = self.client.post(f"/api/reports/{rid}/complete")
        self.assertEqual(blocked.status_code, 409)

    def test_programme_four_without_tare_incomplete(self) -> None:
        iid = self._create_instrument(5.0, 0.5)
        rid = self._create_report(iid)["id"]
        self.client.post(f"/api/reports/{rid}/observations", json={"applied_load": 5000.0, "reference_value": 5000.0, "indication": 5000.5, "zero_error": 0.2})
        self._put_repeatability(rid, "half_max", 15000.0, [15000.0] * 10)
        self._put_repeatability(rid, "near_max", 30000.0, [30000.0] * 10)
        self._put_eccentricity(rid, {"Q1": 0.0, "Q2": 0.0, "Q3": 0.0, "Q4": 0.0})
        self._put_zero_setting(rid, {"setting_mode": "non_automatic", "indication": 0.0})
        prog = self.client.get(f"/api/reports/{rid}").get_json()["test_programme"]
        self.assertEqual(prog["overall_status"], "INCOMPLETE")
        self.assertEqual(prog["incomplete_modules"], ["tare_weighing"])


if __name__ == "__main__":
    unittest.main()
