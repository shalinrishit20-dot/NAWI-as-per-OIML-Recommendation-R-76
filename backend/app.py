import os, uuid, hashlib, math, json
import click

LEGACY_DEMO_SECRET = "nawi-testpro-demo-secret-change-in-production"
_secret_key = os.getenv("SECRET_KEY")
if _secret_key is None or not _secret_key.strip():
    raise RuntimeError("SECRET_KEY must be set to a non-empty value before starting the application")
_secret_key = _secret_key.strip()
if _secret_key == LEGACY_DEMO_SECRET:
    raise RuntimeError("SECRET_KEY must not use the known demo secret value")

from datetime import datetime
from functools import wraps
from flask import Flask, request, jsonify, send_from_directory, send_file, session
from flask_sqlalchemy import SQLAlchemy
from rules import (
    RULESET_VERSION,
    ObservationComplianceResult,
    RepeatabilitySeriesResult,
    EccentricityResult,
    evaluate_digital_observation,
    validate_single_range_classification,
    evaluate_type_approval_repeatability_series,
    evaluate_quarter_segment_eccentricity,
    evaluate_core_test_programme,
    evaluate_zero_setting_accuracy,
    evaluate_tare_weighing_series,
)
from rules.zero_setting import ZeroSettingResult
from rules.tare_weighing import TareWeighingObservationResult, TareWeighingSeriesResult

from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.getenv("NAWI_DATA_DIR") or BASE
FRONTEND_DIST = os.path.join(BASE, "frontend", "dist")
UPLOAD = os.path.join(DATA_DIR, "uploads")
REPORTS = os.path.join(DATA_DIR, "reports")
PRODUCT_IMAGES = os.path.join(UPLOAD, "products")
EVIDENCE = os.path.join(UPLOAD, "evidence")
for p in (UPLOAD, REPORTS, PRODUCT_IMAGES, EVIDENCE): os.makedirs(p, exist_ok=True)

def _demo_mode_enabled():
    return os.getenv("NAWI_DEMO_MODE", "").strip().lower() in {"1", "true", "yes", "on"}


def _secure_session_cookie_enabled():
    return os.getenv("NAWI_SESSION_COOKIE_SECURE", "").strip().lower() not in {"0", "false", "no", "off"}


app = Flask(__name__, static_folder=FRONTEND_DIST, static_url_path="")
app.config["SECRET_KEY"] = _secret_key
app.config["DEMO_MODE"] = _demo_mode_enabled()
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = _secure_session_cookie_enabled()
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///nawi.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 250 * 1024 * 1024

db = SQLAlchemy(app)

ALLOWED_IMAGES = {"png", "jpg", "jpeg", "webp"}
ALLOWED_FILES = ALLOWED_IMAGES | {"pdf", "doc", "docx"}
ALLOWED_VIDEO = {"webm", "mp4", "mov"}

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), default="Tester")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Instrument(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    manufacturer = db.Column(db.String(150), nullable=False)
    model = db.Column(db.String(120), nullable=False)
    serial_number = db.Column(db.String(120), nullable=False, unique=True)
    accuracy_class = db.Column(db.String(20), default="III")
    max_capacity = db.Column(db.Float, nullable=False)
    min_capacity = db.Column(db.Float, nullable=False)
    e = db.Column(db.Float, nullable=False)
    d = db.Column(db.Float, nullable=False)
    unit = db.Column(db.String(20), default="kg")
    software_version = db.Column(db.String(80))
    public_visible = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Report(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_number = db.Column(db.String(60), unique=True, nullable=False)
    instrument_id = db.Column(db.Integer, db.ForeignKey("instrument.id"), nullable=False)
    tester_id = db.Column(db.Integer, db.ForeignKey("user.id"))
    status = db.Column(db.String(30), default="In Progress")
    final_result = db.Column(db.String(25), default="PENDING")
    standard_name = db.Column(db.String(60), default="OIML R 76")
    standard_version = db.Column(db.String(80), default="Configured revision")
    laboratory = db.Column(db.String(150))
    location = db.Column(db.String(150))
    temperature = db.Column(db.Float)
    humidity = db.Column(db.Float)
    pressure = db.Column(db.Float)
    remarks = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    instrument = db.relationship("Instrument", backref="reports")
    tester = db.relationship("User")

class Observation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), nullable=False)
    test_type = db.Column(db.String(80), nullable=False)
    position = db.Column(db.String(80))
    applied_load = db.Column(db.Float, nullable=False)
    reference_value = db.Column(db.Float, nullable=False)
    indication = db.Column(db.Float, nullable=False)
    error = db.Column(db.Float, nullable=False)
    mpe = db.Column(db.Float, nullable=False)
    result = db.Column(db.String(20), nullable=False)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    report = db.relationship("Report", backref="observations")
    calculation = db.relationship(
        "ObservationCalculation",
        backref="observation",
        uselist=False,
        cascade="all, delete-orphan",
    )


class ObservationCalculation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    observation_id = db.Column(db.Integer, db.ForeignKey("observation.id"), unique=True, nullable=False)
    accuracy_class = db.Column(db.String(20), nullable=False)
    verification_stage = db.Column(db.String(20), nullable=False)
    canonical_load = db.Column(db.Float, nullable=False)
    verification_scale_interval = db.Column(db.Float, nullable=False)
    actual_scale_interval = db.Column(db.Float, nullable=False)
    rounding_elimination_required = db.Column(db.Boolean, nullable=False)
    error_method = db.Column(db.String(40), nullable=False)
    additional_load = db.Column(db.Float, nullable=True)
    indication_before_rounding = db.Column(db.Float, nullable=False)
    error_before_rounding = db.Column(db.Float, nullable=False)
    zero_error = db.Column(db.Float, nullable=False)
    corrected_error = db.Column(db.Float, nullable=False)
    absolute_corrected_error = db.Column(db.Float, nullable=False)
    mpe_multiplier = db.Column(db.Float, nullable=False)
    mpe = db.Column(db.Float, nullable=False)
    margin_to_limit = db.Column(db.Float, nullable=False)
    standard_code = db.Column(db.String(40), nullable=False)
    standard_edition = db.Column(db.String(20), nullable=False)
    ruleset_version = db.Column(db.String(40), nullable=False)
    error_rule_reference = db.Column(db.String(120), nullable=False)
    mpe_rule_reference = db.Column(db.String(120), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class RepeatabilityTest(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), nullable=False)
    series_role = db.Column(db.String(20), nullable=False)
    accuracy_class = db.Column(db.String(20), nullable=False)
    max_capacity = db.Column(db.Float, nullable=False)
    applied_load = db.Column(db.Float, nullable=False)
    verification_scale_interval = db.Column(db.Float, nullable=False)
    unit = db.Column(db.String(20), nullable=False)
    readings = db.Column(db.JSON, nullable=False)
    weighing_count = db.Column(db.Integer, nullable=False)
    required_weighings = db.Column(db.Integer, nullable=False)
    minimum_result = db.Column(db.Float, nullable=False)
    maximum_result = db.Column(db.Float, nullable=False)
    spread = db.Column(db.Float, nullable=False)
    mpe = db.Column(db.Float, nullable=False)
    complete = db.Column(db.Boolean, nullable=False)
    spread_within_mpe = db.Column(db.Boolean, nullable=False)
    status = db.Column(db.String(20), nullable=False)
    standard_code = db.Column(db.String(40), nullable=False)
    standard_edition = db.Column(db.String(20), nullable=False)
    ruleset_version = db.Column(db.String(40), nullable=False)
    rule_reference = db.Column(db.String(160), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    report = db.relationship("Report", backref=db.backref("repeatability_tests", cascade="all, delete-orphan"))
    __table_args__ = (db.UniqueConstraint("report_id", "series_role", name="uq_repeatability_report_role"),)


class EccentricityTest(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), unique=True, nullable=False)
    accuracy_class = db.Column(db.String(20), nullable=False)
    max_capacity = db.Column(db.Float, nullable=False)
    verification_scale_interval = db.Column(db.Float, nullable=False)
    unit = db.Column(db.String(20), nullable=False)
    maximum_additive_tare_effect = db.Column(db.Float, nullable=False)
    required_test_load = db.Column(db.Float, nullable=False)
    required_positions = db.Column(db.JSON, nullable=False)
    tested_positions = db.Column(db.JSON, nullable=False)
    position_errors = db.Column(db.JSON, nullable=False)
    failed_positions = db.Column(db.JSON, nullable=False)
    mpe = db.Column(db.Float, nullable=False)
    complete = db.Column(db.Boolean, nullable=False)
    all_positions_within_mpe = db.Column(db.Boolean, nullable=False)
    status = db.Column(db.String(20), nullable=False)
    standard_code = db.Column(db.String(40), nullable=False)
    standard_edition = db.Column(db.String(20), nullable=False)
    ruleset_version = db.Column(db.String(40), nullable=False)
    rule_reference = db.Column(db.String(200), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    report = db.relationship("Report", backref=db.backref("eccentricity_test", uselist=False, cascade="all, delete-orphan"))


class ZeroSettingTest(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), unique=True, nullable=False)
    setting_mode = db.Column(db.String(40), nullable=False)
    indication = db.Column(db.Float, nullable=False)
    additional_load = db.Column(db.Float, nullable=True)
    verification_scale_interval = db.Column(db.Float, nullable=False)
    actual_scale_interval = db.Column(db.Float, nullable=False)
    unit = db.Column(db.String(20), nullable=False)
    rounding_elimination_required = db.Column(db.Boolean, nullable=False)
    error_method = db.Column(db.String(40), nullable=False)
    indication_before_rounding = db.Column(db.Float, nullable=False)
    zero_error = db.Column(db.Float, nullable=False)
    absolute_zero_error = db.Column(db.Float, nullable=False)
    permitted_limit = db.Column(db.Float, nullable=False)
    margin_to_limit = db.Column(db.Float, nullable=False)
    status = db.Column(db.String(20), nullable=False)
    standard_code = db.Column(db.String(40), nullable=False)
    standard_edition = db.Column(db.String(20), nullable=False)
    ruleset_version = db.Column(db.String(40), nullable=False)
    rule_reference = db.Column(db.String(200), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    report = db.relationship("Report", backref=db.backref("zero_setting_test", uselist=False, cascade="all, delete-orphan"))



class TareWeighingTest(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), unique=True, nullable=False)
    tare_value = db.Column(db.Float, nullable=False)
    accuracy_class = db.Column(db.String(20), nullable=False)
    maximum_capacity = db.Column(db.Float, nullable=False)
    minimum_capacity = db.Column(db.Float, nullable=False)
    verification_scale_interval = db.Column(db.Float, nullable=False)
    actual_scale_interval = db.Column(db.Float, nullable=False)
    unit = db.Column(db.String(20), nullable=False)
    required_observations = db.Column(db.Integer, nullable=False)
    observation_count = db.Column(db.Integer, nullable=False)
    complete = db.Column(db.Boolean, nullable=False)
    status = db.Column(db.String(20), nullable=False)
    failed_observation_indices = db.Column(db.JSON, nullable=False)
    observations = db.Column(db.JSON, nullable=False)
    standard_code = db.Column(db.String(40), nullable=False)
    standard_edition = db.Column(db.String(20), nullable=False)
    ruleset_version = db.Column(db.String(40), nullable=False)
    rule_reference = db.Column(db.String(200), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    report = db.relationship("Report", backref=db.backref("tare_weighing_test", uselist=False, cascade="all, delete-orphan"))


class Attachment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), nullable=False)
    filename = db.Column(db.String(255))
    stored_name = db.Column(db.String(255))
    uploaded_at = db.Column(db.DateTime, default=datetime.utcnow)
    report = db.relationship("Report", backref="attachments")

class TestEvidence(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey("report.id"), nullable=False)
    filename = db.Column(db.String(255), nullable=False)
    stored_name = db.Column(db.String(255), nullable=False)
    sha256 = db.Column(db.String(64), nullable=False)
    started_at = db.Column(db.DateTime)
    ended_at = db.Column(db.DateTime)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    report = db.relationship("Report", backref="evidence")

class ProductImage(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    instrument_id = db.Column(db.Integer, db.ForeignKey("instrument.id"), nullable=False)
    filename = db.Column(db.String(255), nullable=False)
    stored_name = db.Column(db.String(255), nullable=False)
    image_hash = db.Column(db.String(64), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    instrument = db.relationship("Instrument", backref="product_images")

class AuditLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"))
    action = db.Column(db.String(255))
    report_id = db.Column(db.Integer)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    user = db.relationship("User")

def current_user():
    uid = session.get("user_id")
    return db.session.get(User, uid) if uid else None

def require_login(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not current_user(): return jsonify({"error":"Authentication required"}), 401
        return fn(*args, **kwargs)
    return wrapper

def require_role(*roles):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            u=current_user()
            if not u: return jsonify({"error":"Authentication required"}),401
            if u.role not in roles: return jsonify({"error":"Permission denied"}),403
            return fn(*args, **kwargs)
        return wrapper
    return deco

def audit(action, report_id=None):
    u=current_user()
    db.session.add(AuditLog(user_id=u.id if u else None, action=action, report_id=report_id))
    db.session.commit()

def image_hash(path):
    # Lightweight perceptual hash for the academic visual-search prototype.
    im=Image.open(path).convert("L").resize((16,16))
    pixels=list(im.getdata()); avg=sum(pixels)/len(pixels)
    bits=''.join('1' if p>=avg else '0' for p in pixels)
    return hex(int(bits,2))[2:].zfill(64)

def hash_distance(a,b):
    try: return bin(int(a,16)^int(b,16)).count('1')
    except: return 999

def public_instrument(i):
    latest=sorted(i.reports, key=lambda r:r.id, reverse=True)
    return {"id":i.id,"manufacturer":i.manufacturer,"model":i.model,"serial_number":i.serial_number,
            "accuracy_class":i.accuracy_class,"max_capacity":i.max_capacity,"min_capacity":i.min_capacity,
            "e":i.e,"d":i.d,"unit":i.unit,"software_version":i.software_version,
            "verification_status":latest[0].final_result if latest else "NOT TESTED",
            "latest_report":latest[0].report_number if latest else None,
            "images":[f"/media/products/{x.stored_name}" for x in i.product_images]}

def calculation_json(c):
    if c is None:
        return None
    return {
        "standard_code": c.standard_code,
        "standard_edition": c.standard_edition,
        "ruleset_version": c.ruleset_version,
        "accuracy_class": c.accuracy_class,
        "verification_stage": c.verification_stage,
        "canonical_load": c.canonical_load,
        "verification_scale_interval": c.verification_scale_interval,
        "actual_scale_interval": c.actual_scale_interval,
        "rounding_elimination_required": c.rounding_elimination_required,
        "error_method": c.error_method,
        "additional_load": c.additional_load,
        "indication_before_rounding": c.indication_before_rounding,
        "error_before_rounding": c.error_before_rounding,
        "zero_error": c.zero_error,
        "corrected_error": c.corrected_error,
        "absolute_corrected_error": c.absolute_corrected_error,
        "mpe_multiplier": c.mpe_multiplier,
        "mpe": c.mpe,
        "margin_to_limit": c.margin_to_limit,
        "error_rule_reference": c.error_rule_reference,
        "mpe_rule_reference": c.mpe_rule_reference,
    }


def _parse_required_number(data, key):
    if key not in data or data[key] is None or data[key] == "":
        raise ValueError(f"{key} is required")
    value = data[key]
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{key} must be a finite number")
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"{key} must be a finite number")
    return number


def _parse_optional_number(data, key):
    if key not in data or data[key] is None or data[key] == "":
        return None
    value = data[key]
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{key} must be a finite number")
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"{key} must be a finite number")
    return number


def _parse_instrument_number(data, key):
    if key not in data or data[key] is None or data[key] == "":
        raise ValueError(f"{key} is required")
    value = data[key]
    if isinstance(value, bool):
        raise ValueError(f"{key} must be a finite number")
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            raise ValueError(f"{key} is required")
        try:
            number = float(text)
        except ValueError:
            raise ValueError(f"{key} must be a finite number") from None
    elif isinstance(value, (int, float)):
        number = float(value)
    else:
        raise ValueError(f"{key} must be a finite number")
    if not math.isfinite(number):
        raise ValueError(f"{key} must be a finite number")
    return number


def _required_identity(data, key):
    if key not in data:
        raise ValueError(f"{key} is required")
    raw = data[key]
    if not isinstance(raw, str):
        raise ValueError(f"{key} must be a non-empty string")
    text = raw.strip()
    if not text:
        raise ValueError(f"{key} is required")
    return text


def classification_json(c):
    return {
        "standard_code": c.standard_code,
        "standard_edition": c.standard_edition,
        "ruleset_version": c.ruleset_version,
        "rule_reference": c.rule_reference,
        "accuracy_class": c.accuracy_class,
        "max_capacity": c.max_capacity,
        "min_capacity": c.min_capacity,
        "verification_scale_interval": c.verification_scale_interval,
        "actual_scale_interval": c.actual_scale_interval,
        "unit": c.unit,
        "number_of_verification_intervals": c.number_of_verification_intervals,
        "minimum_n": c.minimum_n,
        "maximum_n": c.maximum_n,
        "minimum_capacity_multiplier": c.minimum_capacity_multiplier,
        "minimum_capacity_required": c.minimum_capacity_required,
        "minimum_capacity_in_d": c.minimum_capacity_in_d,
        "valid": c.valid,
        "errors": list(c.errors),
    }



def _parse_readings(raw):
    if not isinstance(raw, list):
        raise ValueError("readings must be a non-empty sequence of finite numbers")
    if len(raw) == 0:
        raise ValueError("readings must not be empty")
    out = []
    for i, item in enumerate(raw):
        if isinstance(item, bool) or not isinstance(item, (int, float)):
            raise ValueError(f"readings[{i}] must be a finite number")
        number = float(item)
        if not math.isfinite(number):
            raise ValueError(f"readings[{i}] must be a finite number")
        out.append(number)
    return out


def _parse_position_errors(raw):
    if not isinstance(raw, dict):
        raise ValueError("position_errors must be a mapping of Q1–Q4 to corrected errors")
    if len(raw) == 0:
        raise ValueError("position_errors must not be empty")
    return raw


def _reconstruct_observation_compliance(o):
    c = o.calculation
    if c is None:
        return None
    return ObservationComplianceResult(
        standard_code=c.standard_code,
        standard_edition=c.standard_edition,
        ruleset_version=c.ruleset_version,
        accuracy_class=c.accuracy_class,
        verification_stage=c.verification_stage,
        load=c.canonical_load,
        indication=o.indication,
        verification_scale_interval=c.verification_scale_interval,
        actual_scale_interval=c.actual_scale_interval,
        error_method=c.error_method,
        rounding_elimination_required=c.rounding_elimination_required,
        additional_load=c.additional_load,
        indication_before_rounding=c.indication_before_rounding,
        error_before_rounding=c.error_before_rounding,
        zero_error=c.zero_error,
        corrected_error=c.corrected_error,
        absolute_corrected_error=c.absolute_corrected_error,
        mpe_multiplier=c.mpe_multiplier,
        mpe=c.mpe,
        margin_to_limit=c.margin_to_limit,
        result=o.result,
        error_rule_reference=c.error_rule_reference,
        mpe_rule_reference=c.mpe_rule_reference,
    )


def _reconstruct_repeatability(row):
    if row is None:
        return None
    readings = tuple(float(x) for x in (row.readings or []))
    return RepeatabilitySeriesResult(
        standard_code=row.standard_code,
        standard_edition=row.standard_edition,
        ruleset_version=row.ruleset_version,
        rule_reference=row.rule_reference,
        accuracy_class=row.accuracy_class,
        max_capacity=row.max_capacity,
        applied_load=row.applied_load,
        verification_scale_interval=row.verification_scale_interval,
        unit=row.unit,
        weighing_count=row.weighing_count,
        required_weighings=row.required_weighings,
        minimum_result=row.minimum_result,
        maximum_result=row.maximum_result,
        spread=row.spread,
        mpe=row.mpe,
        complete=row.complete,
        spread_within_mpe=row.spread_within_mpe,
        status=row.status,
        readings=readings,
    )



def _reconstruct_zero_setting(row):
    if row is None:
        return None
    return ZeroSettingResult(
        standard_code=row.standard_code,
        standard_edition=row.standard_edition,
        ruleset_version=row.ruleset_version,
        rule_reference=row.rule_reference,
        setting_mode=row.setting_mode,
        indication=row.indication,
        additional_load=row.additional_load,
        verification_scale_interval=row.verification_scale_interval,
        actual_scale_interval=row.actual_scale_interval,
        unit=row.unit,
        rounding_elimination_required=row.rounding_elimination_required,
        error_method=row.error_method,
        indication_before_rounding=row.indication_before_rounding,
        zero_error=row.zero_error,
        absolute_zero_error=row.absolute_zero_error,
        permitted_limit=row.permitted_limit,
        margin_to_limit=row.margin_to_limit,
        status=row.status,
    )


def _reconstruct_tare_weighing(row):
    if row is None:
        return None
    items = []
    for snap in (row.observations or []):
        method = snap.get("error_method") or "direct"
        compliance = ObservationComplianceResult(
            standard_code=row.standard_code,
            standard_edition=row.standard_edition,
            ruleset_version=snap.get("ruleset_version") or row.ruleset_version,
            accuracy_class=row.accuracy_class,
            verification_stage="initial",
            load=float(snap.get("load") if snap.get("load") is not None else snap.get("net_load") or 0),
            indication=float(snap.get("indication") or 0),
            verification_scale_interval=row.verification_scale_interval,
            actual_scale_interval=row.actual_scale_interval,
            error_method=method,
            rounding_elimination_required=method == "changeover_point",
            additional_load=snap.get("additional_load"),
            indication_before_rounding=float(snap.get("indication_before_rounding") or 0),
            error_before_rounding=float(snap.get("error_before_rounding") or 0),
            zero_error=float(snap.get("zero_error") or 0),
            corrected_error=float(snap.get("corrected_error") or 0),
            absolute_corrected_error=float(snap.get("absolute_corrected_error") or 0),
            mpe_multiplier=1.0,
            mpe=float(snap.get("mpe") or 0),
            margin_to_limit=float(snap.get("margin_to_limit") or 0),
            result=snap.get("result") or "PASS",
            error_rule_reference=snap.get("error_rule_reference") or row.rule_reference,
            mpe_rule_reference=snap.get("mpe_rule_reference") or row.rule_reference,
        )
        items.append(TareWeighingObservationResult(
            net_load=float(snap.get("net_load") or 0),
            gross_load=float(snap.get("gross_load") or 0),
            compliance=compliance,
        ))
    return TareWeighingSeriesResult(
        standard_code=row.standard_code,
        standard_edition=row.standard_edition,
        ruleset_version=row.ruleset_version,
        rule_reference=row.rule_reference,
        accuracy_class=row.accuracy_class,
        maximum_capacity=row.maximum_capacity,
        minimum_capacity=row.minimum_capacity,
        verification_scale_interval=row.verification_scale_interval,
        actual_scale_interval=row.actual_scale_interval,
        unit=row.unit,
        tare_value=row.tare_value,
        required_observations=row.required_observations,
        observation_count=row.observation_count,
        observations=tuple(items),
        failed_observation_indices=tuple(row.failed_observation_indices or ()),
        complete=row.complete,
        status=row.status,
    )


def _reconstruct_eccentricity(row):
    if row is None:
        return None
    raw = row.position_errors or {}
    if isinstance(raw, dict):
        ordered = tuple((k, float(raw[k])) for k in ("Q1", "Q2", "Q3", "Q4") if k in raw)
    else:
        ordered = tuple((p, float(v)) for p, v in raw)
    return EccentricityResult(
        standard_code=row.standard_code,
        standard_edition=row.standard_edition,
        ruleset_version=row.ruleset_version,
        rule_reference=row.rule_reference,
        accuracy_class=row.accuracy_class,
        max_capacity=row.max_capacity,
        maximum_additive_tare_effect=row.maximum_additive_tare_effect,
        verification_scale_interval=row.verification_scale_interval,
        unit=row.unit,
        required_test_load=row.required_test_load,
        required_positions=tuple(row.required_positions or ("Q1", "Q2", "Q3", "Q4")),
        tested_positions=tuple(row.tested_positions or ()),
        mpe=row.mpe,
        position_errors=ordered,
        failed_positions=tuple(row.failed_positions or ()),
        complete=row.complete,
        all_positions_within_mpe=row.all_positions_within_mpe,
        status=row.status,
    )


def repeatability_json(row):
    if row is None:
        return None
    return {
        "id": row.id,
        "series_role": row.series_role,
        "accuracy_class": row.accuracy_class,
        "max_capacity": row.max_capacity,
        "applied_load": row.applied_load,
        "verification_scale_interval": row.verification_scale_interval,
        "unit": row.unit,
        "readings": list(row.readings or []),
        "weighing_count": row.weighing_count,
        "required_weighings": row.required_weighings,
        "minimum_result": row.minimum_result,
        "maximum_result": row.maximum_result,
        "spread": row.spread,
        "mpe": row.mpe,
        "complete": row.complete,
        "spread_within_mpe": row.spread_within_mpe,
        "status": row.status,
        "standard_code": row.standard_code,
        "standard_edition": row.standard_edition,
        "ruleset_version": row.ruleset_version,
        "rule_reference": row.rule_reference,
    }



def tare_weighing_json(row):
    if row is None:
        return None
    return {
        "id": row.id,
        "tare_value": row.tare_value,
        "accuracy_class": row.accuracy_class,
        "maximum_capacity": row.maximum_capacity,
        "minimum_capacity": row.minimum_capacity,
        "verification_scale_interval": row.verification_scale_interval,
        "actual_scale_interval": row.actual_scale_interval,
        "unit": row.unit,
        "required_observations": row.required_observations,
        "observation_count": row.observation_count,
        "complete": row.complete,
        "status": row.status,
        "failed_observation_indices": list(row.failed_observation_indices or []),
        "observations": list(row.observations or []),
        "standard_code": row.standard_code,
        "standard_edition": row.standard_edition,
        "ruleset_version": row.ruleset_version,
        "rule_reference": row.rule_reference,
    }


def zero_setting_json(row):
    if row is None:
        return None
    return {
        "id": row.id,
        "setting_mode": row.setting_mode,
        "indication": row.indication,
        "additional_load": row.additional_load,
        "verification_scale_interval": row.verification_scale_interval,
        "actual_scale_interval": row.actual_scale_interval,
        "unit": row.unit,
        "rounding_elimination_required": row.rounding_elimination_required,
        "error_method": row.error_method,
        "indication_before_rounding": row.indication_before_rounding,
        "zero_error": row.zero_error,
        "absolute_zero_error": row.absolute_zero_error,
        "permitted_limit": row.permitted_limit,
        "margin_to_limit": row.margin_to_limit,
        "status": row.status,
        "standard_code": row.standard_code,
        "standard_edition": row.standard_edition,
        "ruleset_version": row.ruleset_version,
        "rule_reference": row.rule_reference,
    }


def eccentricity_json(row):
    if row is None:
        return None
    return {
        "id": row.id,
        "accuracy_class": row.accuracy_class,
        "max_capacity": row.max_capacity,
        "verification_scale_interval": row.verification_scale_interval,
        "unit": row.unit,
        "maximum_additive_tare_effect": row.maximum_additive_tare_effect,
        "required_test_load": row.required_test_load,
        "required_positions": list(row.required_positions or []),
        "tested_positions": list(row.tested_positions or []),
        "position_errors": dict(row.position_errors) if isinstance(row.position_errors, dict) else {k: v for k, v in (row.position_errors or [])},
        "failed_positions": list(row.failed_positions or []),
        "mpe": row.mpe,
        "complete": row.complete,
        "all_positions_within_mpe": row.all_positions_within_mpe,
        "status": row.status,
        "standard_code": row.standard_code,
        "standard_edition": row.standard_edition,
        "ruleset_version": row.ruleset_version,
        "rule_reference": row.rule_reference,
    }


def programme_json(r):
    reconstructed = []
    legacy = 0
    any_fail = False
    for o in r.observations:
        item = _reconstruct_observation_compliance(o)
        if item is None:
            legacy += 1
        else:
            reconstructed.append(item)
            if item.result == "FAIL":
                any_fail = True
    if legacy and not any_fail:
        observations_for_programme = []
    else:
        observations_for_programme = reconstructed
    half = near = None
    for row in r.repeatability_tests:
        rec = _reconstruct_repeatability(row)
        if row.series_role == "half_max":
            half = rec
        elif row.series_role == "near_max":
            near = rec
    ecc = _reconstruct_eccentricity(r.eccentricity_test)
    zero = _reconstruct_zero_setting(r.zero_setting_test)
    tare = _reconstruct_tare_weighing(r.tare_weighing_test)
    prog = evaluate_core_test_programme(
        observations_for_programme,
        repeatability_half_max=half,
        repeatability_near_max=near,
        eccentricity=ecc,
        zero_setting=zero,
        tare_weighing=tare,
    )
    return {
        "standard_code": prog.standard_code,
        "standard_edition": prog.standard_edition,
        "ruleset_version": prog.ruleset_version,
        "scope_name": prog.scope_name,
        "scope_note": prog.scope_note,
        "full_standard_coverage": False,
        "zero_setting_status": prog.zero_setting_status,
        "observation_status": prog.observation_status,
        "repeatability_status": prog.repeatability_status,
        "eccentricity_status": prog.eccentricity_status,
        "tare_weighing_status": prog.tare_weighing_status,
        "overall_status": prog.overall_status,
        "failed_modules": list(prog.failed_modules),
        "incomplete_modules": list(prog.incomplete_modules),
        "observation_count": len(r.observations),
        "repeatability_half_max_status": prog.repeatability_half_max_status,
        "repeatability_near_max_status": prog.repeatability_near_max_status,
        "rule_references": list(prog.rule_references),
        "legacy_observation_count": legacy,
        "observation_snapshots_complete": legacy == 0,
    }

def report_json(r):
    return {"id":r.id,"report_number":r.report_number,"status":r.status,"final_result":r.final_result,
            "standard_name":r.standard_name,"standard_version":r.standard_version,"laboratory":r.laboratory,
            "location":r.location,"temperature":r.temperature,"humidity":r.humidity,"pressure":r.pressure,
            "remarks":r.remarks,"created_at":r.created_at.isoformat(),"instrument":public_instrument(r.instrument),
            "observations":[{"id":o.id,"test_type":o.test_type,"position":o.position,"applied_load":o.applied_load,
              "reference_value":o.reference_value,"indication":o.indication,"error":o.error,"mpe":o.mpe,"result":o.result,"notes":o.notes,
              "calculation":calculation_json(o.calculation)} for o in r.observations],
            "evidence":[{"id":e.id,"filename":e.filename,"sha256":e.sha256,"started_at":e.started_at.isoformat() if e.started_at else None,
              "ended_at":e.ended_at.isoformat() if e.ended_at else None,"url":f"/media/evidence/{e.stored_name}"} for e in r.evidence],
            "attachments":[{"id":a.id,"filename":a.filename,"url":f"/media/uploads/{a.stored_name}"} for a in r.attachments],
            "repeatability":{
                "half_max": repeatability_json(next((x for x in r.repeatability_tests if x.series_role=="half_max"), None)),
                "near_max": repeatability_json(next((x for x in r.repeatability_tests if x.series_role=="near_max"), None)),
            },
            "eccentricity": eccentricity_json(r.eccentricity_test),
            "zero_setting": zero_setting_json(r.zero_setting_test),
            "tare_weighing": tare_weighing_json(r.tare_weighing_test),
            "test_programme": programme_json(r)}


WORKFLOW_DRAFT = "In Progress"
WORKFLOW_REVIEW = "Review"
WORKFLOW_APPROVED = "Completed"


def _locked_write_response(r):
    return jsonify({
        "error": "Report is locked",
        "detail": "Test data cannot be changed after the report has been submitted for review",
        "status": r.status,
    }), 409


def _assert_report_editable(r):
    if r.status != WORKFLOW_DRAFT:
        return _locked_write_response(r)
    return None


REPORT_READ_ROLES = ("Admin", "Tester", "Reviewer")
REVIEWER_READABLE_STATUSES = (WORKFLOW_REVIEW, WORKFLOW_APPROVED)


def _assert_tester_owns_report(r):
    u = current_user()
    if u.role == "Tester" and r.tester_id != u.id:
        return jsonify({"error": "Permission denied"}), 403
    return None


def _scope_reports_for_current_user(query):
    u = current_user()
    if u.role == "Tester":
        return query.filter(Report.tester_id == u.id)
    if u.role == "Reviewer":
        return query.filter(Report.status.in_(REVIEWER_READABLE_STATUSES))
    return query


def _assert_report_readable(r):
    denied = _assert_tester_owns_report(r)
    if denied:
        return denied
    if current_user().role == "Reviewer" and r.status not in REVIEWER_READABLE_STATUSES:
        return jsonify({
            "error": "Permission denied",
            "detail": "Reviewers may access only reports in Review or Completed status",
            "status": r.status,
        }), 403
    return None

@app.get("/api/health")
def health(): return jsonify({"status":"online","app":"NAWI-TestPro","version":"2.0-react"})

@app.get("/api/auth/me")
def me():
    u=current_user()
    return jsonify({
        "authenticated": bool(u),
        "user": {"id": u.id, "username": u.username, "role": u.role} if u else None,
        "demo_mode": app.config["DEMO_MODE"],
    })

@app.post("/api/auth/login")
def login():
    data=request.get_json() or {}
    u=User.query.filter_by(username=(data.get("username") or "").strip()).first()
    if not u or not check_password_hash(u.password_hash,data.get("password") or ""):
        return jsonify({"error":"Invalid credentials"}),401
    session["user_id"]=u.id; audit("Logged in")
    return jsonify({"user":{"id":u.id,"username":u.username,"role":u.role}})

@app.post("/api/auth/logout")
def logout():
    if current_user(): audit("Logged out")
    session.clear(); return jsonify({"ok":True})

@app.get("/api/dashboard")
@require_login
def dashboard_api():
    u=current_user()
    report_query=Report.query if u.role == "Manufacturer" else _scope_reports_for_current_user(Report.query)
    recent_query=Report.query.order_by(Report.id.desc())
    reports=[] if u.role == "Manufacturer" else _scope_reports_for_current_user(recent_query).all()
    return jsonify({"stats":{"instruments":Instrument.query.count(),"reports":report_query.count(),
      "in_progress":report_query.filter_by(status="In Progress").count(),"review":report_query.filter_by(status="Review").count(),
      "completed":report_query.filter_by(status="Completed").count(),"passed":report_query.filter_by(final_result="PASS").count(),
      "failed":report_query.filter_by(final_result="FAIL").count()},"recent":[report_json(r) for r in reports[:8]]})

@app.get("/api/instruments")
@require_login
def instruments_api(): return jsonify([public_instrument(i) for i in Instrument.query.order_by(Instrument.id.desc()).all()])

@app.post("/api/instruments")
@require_role("Admin","Tester")
def instrument_create():
    d=request.get_json() or {}
    try:
        manufacturer=_required_identity(d,"manufacturer")
        model=_required_identity(d,"model")
        serial_number=_required_identity(d,"serial_number")
        accuracy_class=d.get("accuracy_class", "III")
        if isinstance(accuracy_class, str):
            accuracy_class=accuracy_class.strip()
        max_capacity=_parse_instrument_number(d,"max_capacity")
        min_capacity=_parse_instrument_number(d,"min_capacity")
        e=_parse_instrument_number(d,"e")
        d_interval=_parse_instrument_number(d,"d")
        unit=d.get("unit","kg")
        if isinstance(unit, str):
            unit=unit.strip()
        classification=validate_single_range_classification(
            accuracy_class, max_capacity, min_capacity, e, d_interval, unit,
        )
        if not classification.valid:
            return jsonify({
                "error":"Instrument does not satisfy declared OIML R76 classification",
                "detail":" ".join(classification.errors),
                "classification":classification_json(classification),
            }),422
        i=Instrument(
            manufacturer=manufacturer, model=model, serial_number=serial_number,
            accuracy_class=classification.accuracy_class,
            max_capacity=classification.max_capacity,
            min_capacity=classification.min_capacity,
            e=classification.verification_scale_interval,
            d=classification.actual_scale_interval,
            unit=classification.unit,
            software_version=d.get("software_version"),
            public_visible=bool(d.get("public_visible", True)),
        )
        db.session.add(i)
        db.session.flush()
        u=current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Registered instrument {i.id} [{RULESET_VERSION} classification valid]",
            report_id=None,
        ))
        db.session.commit()
        body=public_instrument(i)
        body["classification"]=classification_json(classification)
        return jsonify(body),201
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error":"Invalid instrument data","detail":str(ex)}),400
    except Exception:
        db.session.rollback()
        return jsonify({"error":"Could not register instrument"}),500

@app.get("/api/reports")
@require_role(*REPORT_READ_ROLES)
def reports_api():
    q=(request.args.get("q") or "").strip()
    query=Report.query.join(Instrument)
    if q:
        query=query.filter((Report.report_number.contains(q))|(Instrument.manufacturer.contains(q))|(Instrument.model.contains(q))|(Instrument.serial_number.contains(q)))
    query=_scope_reports_for_current_user(query)
    return jsonify([report_json(r) for r in query.order_by(Report.id.desc()).all()])

@app.post("/api/reports")
@require_role("Admin","Tester")
def report_create():
    d=request.get_json() or {}; i=db.session.get(Instrument,d.get("instrument_id"))
    if not i: return jsonify({"error":"Instrument not found"}),404
    no="NAWI-"+datetime.now().strftime("%Y%m%d-%H%M%S")+"-"+uuid.uuid4().hex[:4].upper()
    r=Report(report_number=no,instrument_id=i.id,tester_id=current_user().id,standard_version=RULESET_VERSION,
             laboratory=d.get("laboratory"),location=d.get("location"))
    db.session.add(r); db.session.commit(); audit(f"Created report {no}",r.id)
    return jsonify(report_json(r)),201

@app.get("/api/reports/<int:rid>")
@require_role(*REPORT_READ_ROLES)
def report_get(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Not found"}),404
    denied=_assert_report_readable(r)
    if denied: return denied
    return jsonify(report_json(r))

@app.post("/api/reports/<int:rid>/observations")
@require_role("Admin","Tester")
def observation_create(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    denied = _assert_tester_owns_report(r)
    if denied: return denied
    locked = _assert_report_editable(r)
    if locked:
        return locked
    d=request.get_json() or {}
    try:
        load=_parse_required_number(d,"applied_load")
        if load < 0:
            raise ValueError("applied_load must be >= 0")
        ref=_parse_required_number(d,"reference_value")
        ind=_parse_required_number(d,"indication")
        zero_error=_parse_required_number(d,"zero_error")
        additional_load=_parse_optional_number(d,"additional_load")
        stage=d.get("verification_stage") or "initial"
        inst=r.instrument
        # Canonical L for both error determination and Table 6 MPE is reference_value.
        # applied_load is stored as the original test-loading record only.
        compliance=evaluate_digital_observation(
            accuracy_class=inst.accuracy_class,
            load=ref,
            indication=ind,
            e=inst.e,
            d=inst.d,
            zero_error=zero_error,
            stage=stage,
            additional_load=additional_load,
        )
        o=Observation(
            report_id=rid,
            test_type=d.get("test_type","Accuracy"),
            position=d.get("position"),
            applied_load=load,
            reference_value=ref,
            indication=ind,
            error=compliance.corrected_error,
            mpe=compliance.mpe,
            result=compliance.result,
            notes=d.get("notes"),
        )
        db.session.add(o)
        db.session.flush()
        db.session.add(ObservationCalculation(
            observation_id=o.id,
            accuracy_class=compliance.accuracy_class,
            verification_stage=compliance.verification_stage,
            canonical_load=compliance.load,
            verification_scale_interval=compliance.verification_scale_interval,
            actual_scale_interval=compliance.actual_scale_interval,
            rounding_elimination_required=compliance.rounding_elimination_required,
            error_method=compliance.error_method,
            additional_load=compliance.additional_load,
            indication_before_rounding=compliance.indication_before_rounding,
            error_before_rounding=compliance.error_before_rounding,
            zero_error=compliance.zero_error,
            corrected_error=compliance.corrected_error,
            absolute_corrected_error=compliance.absolute_corrected_error,
            mpe_multiplier=compliance.mpe_multiplier,
            mpe=compliance.mpe,
            margin_to_limit=compliance.margin_to_limit,
            standard_code=compliance.standard_code,
            standard_edition=compliance.standard_edition,
            ruleset_version=compliance.ruleset_version,
            error_rule_reference=compliance.error_rule_reference,
            mpe_rule_reference=compliance.mpe_rule_reference,
        ))
        u = current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Added {o.test_type} observation: {o.result} [{RULESET_VERSION}]",
            report_id=rid,
        ))
        db.session.commit()
        return jsonify(report_json(r)),201
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error":"Invalid observation","detail":str(ex)}),400
    except Exception:
        db.session.rollback()
        return jsonify({"error":"Could not save observation"}),500


@app.put("/api/reports/<int:rid>/repeatability/<series_role>")
@require_role("Admin","Tester")
def repeatability_put(rid, series_role):
    r = db.session.get(Report, rid)
    if not r:
        return jsonify({"error": "Report not found"}), 404
    denied = _assert_tester_owns_report(r)
    if denied:
        return denied
    if series_role not in ("half_max", "near_max"):
        return jsonify({"error": "Invalid repeatability series role"}), 400
    locked = _assert_report_editable(r)
    if locked:
        return locked
    d = request.get_json() or {}
    try:
        applied_load = _parse_required_number(d, "applied_load")
        readings = _parse_readings(d.get("readings"))
        inst = r.instrument
        result = evaluate_type_approval_repeatability_series(
            inst.accuracy_class, inst.max_capacity, applied_load, inst.e, inst.unit, readings,
        )
        row = RepeatabilityTest.query.filter_by(report_id=rid, series_role=series_role).first()
        if row is None:
            row = RepeatabilityTest(report_id=rid, series_role=series_role)
            db.session.add(row)
        row.accuracy_class = result.accuracy_class
        row.max_capacity = result.max_capacity
        row.applied_load = result.applied_load
        row.verification_scale_interval = result.verification_scale_interval
        row.unit = result.unit
        row.readings = list(result.readings)
        row.weighing_count = result.weighing_count
        row.required_weighings = result.required_weighings
        row.minimum_result = result.minimum_result
        row.maximum_result = result.maximum_result
        row.spread = result.spread
        row.mpe = result.mpe
        row.complete = result.complete
        row.spread_within_mpe = result.spread_within_mpe
        row.status = result.status
        row.standard_code = result.standard_code
        row.standard_edition = result.standard_edition
        row.ruleset_version = result.ruleset_version
        row.rule_reference = result.rule_reference
        row.updated_at = datetime.utcnow()
        u = current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Evaluated repeatability {series_role} [{RULESET_VERSION}: {result.status}]",
            report_id=rid,
        ))
        db.session.commit()
        return jsonify(repeatability_json(row)), 200
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error": "Invalid repeatability data", "detail": str(ex)}), 400
    except Exception:
        db.session.rollback()
        return jsonify({"error": "Could not save repeatability test"}), 500


@app.put("/api/reports/<int:rid>/eccentricity")
@require_role("Admin","Tester")
def eccentricity_put(rid):
    r = db.session.get(Report, rid)
    if not r:
        return jsonify({"error": "Report not found"}), 404
    denied = _assert_tester_owns_report(r)
    if denied:
        return denied
    locked = _assert_report_editable(r)
    if locked:
        return locked
    d = request.get_json() or {}
    try:
        if "maximum_additive_tare_effect" not in d or d.get("maximum_additive_tare_effect") in (None, ""):
            tare = 0.0
        else:
            tare = _parse_required_number(d, "maximum_additive_tare_effect")
        positions = _parse_position_errors(d.get("position_errors"))
        inst = r.instrument
        result = evaluate_quarter_segment_eccentricity(
            inst.accuracy_class, inst.max_capacity, inst.e, inst.unit, positions,
            maximum_additive_tare_effect=tare,
        )
        row = r.eccentricity_test
        if row is None:
            row = EccentricityTest(report_id=rid)
            db.session.add(row)
        row.accuracy_class = result.accuracy_class
        row.max_capacity = result.max_capacity
        row.verification_scale_interval = result.verification_scale_interval
        row.unit = result.unit
        row.maximum_additive_tare_effect = result.maximum_additive_tare_effect
        row.required_test_load = result.required_test_load
        row.required_positions = list(result.required_positions)
        row.tested_positions = list(result.tested_positions)
        row.position_errors = {k: v for k, v in result.position_errors}
        row.failed_positions = list(result.failed_positions)
        row.mpe = result.mpe
        row.complete = result.complete
        row.all_positions_within_mpe = result.all_positions_within_mpe
        row.status = result.status
        row.standard_code = result.standard_code
        row.standard_edition = result.standard_edition
        row.ruleset_version = result.ruleset_version
        row.rule_reference = result.rule_reference
        row.updated_at = datetime.utcnow()
        u = current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Evaluated eccentricity [{RULESET_VERSION}: {result.status}]",
            report_id=rid,
        ))
        db.session.commit()
        return jsonify(eccentricity_json(row)), 200
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error": "Invalid eccentricity data", "detail": str(ex)}), 400
    except Exception:
        db.session.rollback()
        return jsonify({"error": "Could not save eccentricity test"}), 500


@app.put("/api/reports/<int:rid>/zero-setting")
@require_role("Admin","Tester")
def zero_setting_put(rid):
    r = db.session.get(Report, rid)
    if not r:
        return jsonify({"error": "Report not found"}), 404
    denied = _assert_tester_owns_report(r)
    if denied:
        return denied
    locked = _assert_report_editable(r)
    if locked:
        return locked
    d = request.get_json() or {}
    try:
        mode = d.get("setting_mode", "non_automatic")
        indication = _parse_required_number(d, "indication")
        additional_load = _parse_optional_number(d, "additional_load")
        inst = r.instrument
        result = evaluate_zero_setting_accuracy(
            indication=indication,
            e=inst.e,
            d=inst.d,
            unit=inst.unit,
            additional_load=additional_load,
            setting_mode=mode,
        )
        row = r.zero_setting_test
        if row is None:
            row = ZeroSettingTest(report_id=rid)
            db.session.add(row)
        row.setting_mode = result.setting_mode
        row.indication = result.indication
        row.additional_load = result.additional_load
        row.verification_scale_interval = result.verification_scale_interval
        row.actual_scale_interval = result.actual_scale_interval
        row.unit = result.unit
        row.rounding_elimination_required = result.rounding_elimination_required
        row.error_method = result.error_method
        row.indication_before_rounding = result.indication_before_rounding
        row.zero_error = result.zero_error
        row.absolute_zero_error = result.absolute_zero_error
        row.permitted_limit = result.permitted_limit
        row.margin_to_limit = result.margin_to_limit
        row.status = result.status
        row.standard_code = result.standard_code
        row.standard_edition = result.standard_edition
        row.ruleset_version = result.ruleset_version
        row.rule_reference = result.rule_reference
        row.updated_at = datetime.utcnow()
        u = current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Saved zero-setting accuracy: {result.status}",
            report_id=rid,
        ))
        db.session.commit()
        return jsonify(zero_setting_json(row)), 200
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error": "Invalid zero-setting data", "detail": str(ex)}), 400
    except Exception:
        db.session.rollback()
        return jsonify({"error": "Could not save zero-setting test"}), 500


@app.put("/api/reports/<int:rid>/tare-weighing")
@require_role("Admin","Tester")
def tare_weighing_put(rid):
    r = db.session.get(Report, rid)
    if not r:
        return jsonify({"error": "Report not found"}), 404
    denied = _assert_tester_owns_report(r)
    if denied:
        return denied
    locked = _assert_report_editable(r)
    if locked:
        return locked
    d = request.get_json() or {}
    try:
        tare_value = _parse_required_number(d, "tare_value")
        observations = d.get("observations")
        if not isinstance(observations, list):
            raise ValueError("observations must be a sequence of mappings")
        inst = r.instrument
        result = evaluate_tare_weighing_series(
            inst.accuracy_class,
            inst.max_capacity,
            inst.min_capacity,
            inst.e,
            inst.d,
            inst.unit,
            tare_value,
            observations,
        )
        payload = []
        for item in result.observations:
            c = item.compliance
            payload.append({
                "net_load": item.net_load,
                "gross_load": item.gross_load,
                "load": c.load,
                "indication": c.indication,
                "additional_load": c.additional_load,
                "indication_before_rounding": c.indication_before_rounding,
                "error_before_rounding": c.error_before_rounding,
                "zero_error": c.zero_error,
                "corrected_error": c.corrected_error,
                "absolute_corrected_error": c.absolute_corrected_error,
                "mpe": c.mpe,
                "margin_to_limit": c.margin_to_limit,
                "result": c.result,
                "error_method": c.error_method,
                "ruleset_version": c.ruleset_version,
                "error_rule_reference": c.error_rule_reference,
                "mpe_rule_reference": c.mpe_rule_reference,
            })
        row = r.tare_weighing_test
        if row is None:
            row = TareWeighingTest(report_id=rid)
            db.session.add(row)
        row.tare_value = result.tare_value
        row.accuracy_class = result.accuracy_class
        row.maximum_capacity = result.maximum_capacity
        row.minimum_capacity = result.minimum_capacity
        row.verification_scale_interval = result.verification_scale_interval
        row.actual_scale_interval = result.actual_scale_interval
        row.unit = result.unit
        row.required_observations = result.required_observations
        row.observation_count = result.observation_count
        row.complete = result.complete
        row.status = result.status
        row.failed_observation_indices = list(result.failed_observation_indices)
        row.observations = payload
        row.standard_code = result.standard_code
        row.standard_edition = result.standard_edition
        row.ruleset_version = result.ruleset_version
        row.rule_reference = result.rule_reference
        row.updated_at = datetime.utcnow()
        u = current_user()
        db.session.add(AuditLog(
            user_id=u.id if u else None,
            action=f"Saved tare weighing: {result.status}",
            report_id=rid,
        ))
        db.session.commit()
        return jsonify(tare_weighing_json(row)), 200
    except ValueError as ex:
        db.session.rollback()
        return jsonify({"error": "Invalid tare weighing data", "detail": str(ex)}), 400
    except Exception:
        db.session.rollback()
        return jsonify({"error": "Could not save tare weighing test"}), 500

@app.post("/api/reports/<int:rid>/environment")
@require_role("Admin","Tester")
def environment(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    denied = _assert_tester_owns_report(r)
    if denied: return denied
    locked = _assert_report_editable(r)
    if locked: return locked
    d=request.get_json() or {}
    for key in ("temperature","humidity","pressure"):
        setattr(r,key,float(d[key]) if d.get(key) not in (None,"") else None)
    r.laboratory=d.get("laboratory"); r.location=d.get("location"); r.remarks=d.get("remarks")
    db.session.commit(); audit("Updated environmental conditions",rid)
    return jsonify(report_json(r))

@app.post("/api/reports/<int:rid>/complete")
@require_role("Admin","Tester")
def complete(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    denied = _assert_tester_owns_report(r)
    if denied: return denied
    if r.status == WORKFLOW_REVIEW:
        return jsonify({"error":"Report already completed","status":r.status}),409
    if r.status == WORKFLOW_APPROVED:
        return jsonify({"error":"Approved reports cannot be completed again","status":r.status}),409
    if r.status != WORKFLOW_DRAFT:
        return jsonify({"error":"Invalid workflow transition","status":r.status}),409
    prog = programme_json(r)
    if prog["overall_status"] == "INCOMPLETE":
        return jsonify({
            "error": "Implemented core programme is incomplete",
            "code": "programme_incomplete",
            "overall_status": prog["overall_status"],
            "incomplete_modules": prog["incomplete_modules"],
            "scope_note": prog["scope_note"],
        }), 409
    if prog["overall_status"] not in ("PASS", "FAIL"):
        return jsonify({"error":"Cannot complete report","overall_status":prog["overall_status"]}),409
    r.final_result = prog["overall_status"]
    r.status = WORKFLOW_REVIEW
    u = current_user()
    db.session.add(AuditLog(
        user_id=u.id if u else None,
        action=f"Submitted for review: {r.final_result}",
        report_id=rid,
    ))
    db.session.commit()
    return jsonify(report_json(r))

@app.post("/api/reports/<int:rid>/approve")
@require_role("Admin","Reviewer")
def approve(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    if r.status == WORKFLOW_APPROVED:
        return jsonify({"error":"Report already approved","status":r.status}),409
    if r.status != WORKFLOW_REVIEW:
        return jsonify({
            "error": "Report must be completed before approval",
            "status": r.status,
        }), 409
    r.status = WORKFLOW_APPROVED
    u = current_user()
    db.session.add(AuditLog(
        user_id=u.id if u else None,
        action="Report approved",
        report_id=rid,
    ))
    db.session.commit()
    return jsonify(report_json(r))

@app.post("/api/reports/<int:rid>/upload")
@require_role("Admin","Tester")
def upload_attachment(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    denied = _assert_tester_owns_report(r)
    if denied: return denied
    locked = _assert_report_editable(r)
    if locked: return locked
    f=request.files.get("file")
    if not f or not f.filename: return jsonify({"error":"File missing"}),400
    ext=f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    if ext not in ALLOWED_FILES: return jsonify({"error":"File type not allowed"}),400
    stored=uuid.uuid4().hex+"_"+secure_filename(f.filename); f.save(os.path.join(UPLOAD,stored))
    db.session.add(Attachment(report_id=rid,filename=secure_filename(f.filename),stored_name=stored)); db.session.commit(); audit("Uploaded attachment",rid)
    return jsonify(report_json(r))

@app.post("/api/reports/<int:rid>/evidence")
@require_role("Admin","Tester")
def evidence_upload(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Report not found"}),404
    denied = _assert_tester_owns_report(r)
    if denied: return denied
    locked = _assert_report_editable(r)
    if locked: return locked
    f=request.files.get("video")
    if not f: return jsonify({"error":"Video missing"}),400
    ext=f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    if ext not in ALLOWED_VIDEO: return jsonify({"error":"Video type not allowed"}),400
    stored=uuid.uuid4().hex+"_evidence."+ext; path=os.path.join(EVIDENCE,stored); f.save(path)
    sha=hashlib.sha256(open(path,"rb").read()).hexdigest()
    def dt(v):
        try:return datetime.fromisoformat(v.replace("Z","+00:00")).replace(tzinfo=None) if v else None
        except:return None
    e=TestEvidence(report_id=rid,filename=secure_filename(f.filename),stored_name=stored,sha256=sha,started_at=dt(request.form.get("started_at")),ended_at=dt(request.form.get("ended_at")))
    db.session.add(e); db.session.commit(); audit("Stored live test video evidence",rid)
    return jsonify(report_json(r))

@app.post("/api/instruments/<int:iid>/images")
@require_role("Admin","Manufacturer")
def product_images(iid):
    i=db.session.get(Instrument,iid); files=request.files.getlist("images")
    if not i: return jsonify({"error":"Instrument not found"}),404
    added=0
    for f in files:
        if not f or not f.filename: continue
        ext=f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
        if ext not in ALLOWED_IMAGES: continue
        stored=uuid.uuid4().hex+"."+ext; path=os.path.join(PRODUCT_IMAGES,stored); f.save(path)
        try: h=image_hash(path)
        except: os.remove(path); continue
        db.session.add(ProductImage(instrument_id=iid,filename=secure_filename(f.filename),stored_name=stored,image_hash=h)); added+=1
    db.session.commit(); audit(f"Added {added} product image(s) to instrument {iid}")
    return jsonify(public_instrument(i))

@app.post("/api/visual-search")
def visual_search():
    f=request.files.get("image")
    if not f: return jsonify({"error":"Image required"}),400
    ext=f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else "jpg"
    if ext not in ALLOWED_IMAGES: return jsonify({"error":"Image type not supported"}),400
    temp=os.path.join(UPLOAD,"search_"+uuid.uuid4().hex+"."+ext); f.save(temp)
    try: target=image_hash(temp)
    finally:
        try: os.remove(temp)
        except: pass
    results=[]
    for p in ProductImage.query.all():
        dist=hash_distance(target,p.image_hash); confidence=max(0,round(100-(dist/256*100),1))
        results.append((confidence,p.instrument))
    results.sort(key=lambda x:x[0],reverse=True)
    unique=[]; seen=set()
    for confidence,i in results:
        if i.id in seen: continue
        seen.add(i.id); unique.append({"confidence":confidence,"instrument":public_instrument(i)})
        if len(unique)>=5: break
    return jsonify({"results":unique,"note":"AI-assisted visual identification prototype. Use serial/QR verification for definitive physical identity."})

@app.get("/api/public/instruments/<int:iid>")
def public_instrument_api(iid):
    i=db.session.get(Instrument,iid)
    if not i or not i.public_visible: return jsonify({"error":"Not found"}),404
    return jsonify(public_instrument(i))

@app.get("/api/verify/<report_number>")
def verify(report_number):
    r=Report.query.filter_by(report_number=report_number).first()
    if not r: return jsonify({"valid":False,"error":"Report not found"}),404
    evidence_hashes=[e.sha256 for e in r.evidence]
    return jsonify({"valid":True,"report_number":r.report_number,"instrument":public_instrument(r.instrument),"status":r.status,"final_result":r.final_result,
      "test_date":r.created_at.isoformat(),"evidence_count":len(r.evidence),"evidence_integrity":bool(evidence_hashes),"evidence_hashes":evidence_hashes})

@app.get("/api/audit")
@require_role("Admin")
def audit_api():
    return jsonify([{"id":x.id,"action":x.action,"user":x.user.username if x.user else "system","report_id":x.report_id,"created_at":x.created_at.isoformat()} for x in AuditLog.query.order_by(AuditLog.id.desc()).limit(200).all()])


def _fmt(v):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def _build_report_document_data(report):
    """Normalized document payload from persisted snapshots only."""
    inst = report.instrument
    prog = programme_json(report)
    half = repeatability_json(next((x for x in report.repeatability_tests if x.series_role == "half_max"), None))
    near = repeatability_json(next((x for x in report.repeatability_tests if x.series_role == "near_max"), None))
    ecc = eccentricity_json(report.eccentricity_test)
    zero_setting = zero_setting_json(report.zero_setting_test)
    tare_weighing = tare_weighing_json(report.tare_weighing_test)
    observations = []
    for o in report.observations:
        calc = calculation_json(o.calculation)
        observations.append({
            "test_type": o.test_type,
            "position": o.position,
            "applied_load": o.applied_load,
            "reference_value": o.reference_value,
            "indication": o.indication,
            "error": o.error,
            "mpe": o.mpe,
            "result": o.result,
            "legacy": calc is None,
            "calculation": calc,
        })
    return {
        "title": "OIML R 76 Implemented Core Test Report",
        "subtitle": "R 76-2:2007 aligned report — implemented test scope",
        "scope_warning": (
            "Overall PASS indicates only that all currently implemented core modules passed. "
            "This is not a complete OIML R 76 type-evaluation or type-approval verdict."
        ),
        "format_basis": "OIML R 76-2:2007 — implemented subset",
        "report_number": report.report_number,
        "report_id": report.id,
        "status": report.status,
        "final_result": report.final_result,
        "created_at": report.created_at.isoformat() if report.created_at else None,
        "ruleset_version": report.standard_version or RULESET_VERSION,
        "instrument": {
            "id": inst.id,
            "manufacturer": inst.manufacturer,
            "model": inst.model,
            "serial_number": inst.serial_number,
            "accuracy_class": inst.accuracy_class,
            "max_capacity": inst.max_capacity,
            "min_capacity": inst.min_capacity,
            "e": inst.e,
            "d": inst.d,
            "n": "Not recorded",
            "unit": inst.unit,
            "software_version": inst.software_version or "Not recorded",
        },
        "programme": prog,
        "observations": observations,
        "repeatability": {"half_max": half, "near_max": near},
        "eccentricity": ecc,
        "zero_setting": zero_setting,
        "tare_weighing": tare_weighing,
    }


def _pdf_kv_table(rows, styles):
    data = [[Paragraph(f"<b>{a}</b>", styles["Normal"]), Paragraph(_fmt(b), styles["Normal"])] for a, b in rows]
    tbl = Table(data, colWidths=[160, 350])
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return tbl


def _pdf_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Times-Roman", 8)
    canvas.drawString(28, 16, "OIML R 76 implemented core test report")
    canvas.drawRightString(A4[0] - 28, 16, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(r, path):
    data = _build_report_document_data(r)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Wrap", fontName="Times-Roman", fontSize=8, leading=10))
    story = [
        Paragraph(data["title"], styles["Title"]),
        Paragraph(data["subtitle"], styles["Italic"]),
        Spacer(1, 8),
        Paragraph(data["programme"]["scope_note"], styles["Italic"]),
        Paragraph(data["scope_warning"], styles["Italic"]),
        Spacer(1, 10),
        Paragraph("General information concerning the type", styles["Heading2"]),
        _pdf_kv_table([
            ("Report ID", data["report_id"]),
            ("Report number", data["report_number"]),
            ("Instrument ID", data["instrument"]["id"]),
            ("Manufacturer", data["instrument"]["manufacturer"]),
            ("Type designation / model", data["instrument"]["model"]),
            ("Serial number", data["instrument"]["serial_number"]),
            ("Accuracy class", data["instrument"]["accuracy_class"]),
            ("Max", f"{_fmt(data['instrument']['max_capacity'])} {data['instrument']['unit']}"),
            ("Min", f"{_fmt(data['instrument']['min_capacity'])} {data['instrument']['unit']}"),
            ("Verification scale interval e", f"{_fmt(data['instrument']['e'])} {data['instrument']['unit']}"),
            ("Actual scale interval d", f"{_fmt(data['instrument']['d'])} {data['instrument']['unit']}"),
            ("Number of verification intervals n", _fmt(data["instrument"]["n"])),
            ("Software version", data["instrument"]["software_version"]),
            ("Report status", data["status"]),
            ("Created", data["created_at"]),
            ("Ruleset version", data["ruleset_version"]),
        ], styles),
        Spacer(1, 10),
        Paragraph("Summary of Implemented Test Evaluation", styles["Heading2"]),
        Paragraph(
            "Overall PASS indicates only that all currently implemented core modules passed.",
            styles["Italic"],
        ),
        _pdf_kv_table([
            ("Overall implemented-core status", data["programme"]["overall_status"]),
            ("Zero-setting", data["programme"].get("zero_setting_status") or "INCOMPLETE"),
            ("Observation / weighing checks", data["programme"]["observation_status"]),
            ("Repeatability", data["programme"]["repeatability_status"]),
            ("Eccentricity", data["programme"]["eccentricity_status"]),
            ("Tare weighing", data["programme"].get("tare_weighing_status") or "INCOMPLETE"),
            ("Repeatability Half-Max", data["programme"]["repeatability_half_max_status"] or "Not recorded"),
            ("Repeatability Near-Max", data["programme"]["repeatability_near_max_status"] or "Not recorded"),
            ("Failed modules", ", ".join(data["programme"]["failed_modules"]) or "none"),
            ("Incomplete modules", ", ".join(data["programme"]["incomplete_modules"]) or "none"),
            ("Ruleset", data["programme"]["ruleset_version"]),
            ("full_standard_coverage", "No"),
            ("Scope note", data["programme"]["scope_note"]),
        ], styles),
        Spacer(1, 10),
        Paragraph("Zero-setting accuracy", styles["Heading2"]),
    ]
    zs = data.get("zero_setting")
    if not zs:
        story.append(Paragraph("Not recorded — test not performed", styles["Normal"]))
    else:
        story.append(_pdf_kv_table([
            ("Status", zs["status"]),
            ("Setting mode", zs["setting_mode"]),
            ("Indication", f"{_fmt(zs['indication'])} {zs['unit']}"),
            ("Additional load", "—" if zs["additional_load"] is None else f"{_fmt(zs['additional_load'])} {zs['unit']}"),
            ("Error method", zs["error_method"]),
            ("Indication before rounding", f"{_fmt(zs['indication_before_rounding'])} {zs['unit']}"),
            ("Zero error", f"{_fmt(zs['zero_error'])} {zs['unit']}"),
            ("Absolute zero error", f"{_fmt(zs['absolute_zero_error'])} {zs['unit']}"),
            ("Permitted limit", f"{_fmt(zs['permitted_limit'])} {zs['unit']}"),
            ("Margin to limit", f"{_fmt(zs['margin_to_limit'])} {zs['unit']}"),
            ("e", f"{_fmt(zs['verification_scale_interval'])} {zs['unit']}"),
            ("d", f"{_fmt(zs['actual_scale_interval'])} {zs['unit']}"),
            ("Ruleset", zs["ruleset_version"]),
            ("Rule reference", zs["rule_reference"]),
        ], styles))
    story += [
        Spacer(1, 10),
        Paragraph("Section 1: Weighing performance", styles["Heading2"]),
    ]
    headers = ["Test", "Pos", "Applied", "L", "I", "ΔL", "P", "E", "E0", "Ec", "mpe", "Result"]
    table = [headers]
    for o in data["observations"]:
        c = o["calculation"]
        if o["legacy"] or c is None:
            table.append([
                Paragraph(o["test_type"] or "—", styles["Wrap"]),
                o["position"] or "—",
                _fmt(o["applied_load"]),
                _fmt(o["reference_value"]),
                _fmt(o["indication"]),
                "—", "—", "—", "—", "—", "—",
                Paragraph("Legacy / no OIML calculation snapshot", styles["Wrap"]),
            ])
        else:
            table.append([
                Paragraph(str(o["test_type"] or "—"), styles["Wrap"]),
                o["position"] or "—",
                _fmt(o["applied_load"]),
                _fmt(c["canonical_load"]),
                _fmt(o["indication"]),
                _fmt(c["additional_load"]),
                _fmt(c["indication_before_rounding"]),
                _fmt(c["error_before_rounding"]),
                _fmt(c["zero_error"]),
                _fmt(c["corrected_error"]),
                _fmt(c["mpe"]),
                o["result"],
            ])
    if len(table) == 1:
        table.append(["No observations recorded"] + [""] * 11)
    widths = [48, 32, 42, 42, 42, 32, 38, 38, 32, 38, 38, 70]
    obs_tbl = Table(table, colWidths=widths, repeatRows=1)
    obs_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [obs_tbl, Spacer(1, 8)]
    story.append(Paragraph("Observation calculation snapshots", styles["Heading3"] if "Heading3" in styles else styles["Heading2"]))
    for i, o in enumerate(data["observations"], 1):
        c = o["calculation"]
        if o["legacy"] or c is None:
            story.append(Paragraph(f"Observation {i}: Legacy / no OIML calculation snapshot", styles["Normal"]))
            continue
        story.append(_pdf_kv_table([
            (f"Observation {i} result", o["result"]),
            ("|Ec|", c["absolute_corrected_error"]),
            ("Margin to limit", c["margin_to_limit"]),
            ("Error method", c["error_method"]),
            ("Ruleset", c["ruleset_version"]),
            ("Error rule reference", c["error_rule_reference"]),
            ("MPE rule reference", c["mpe_rule_reference"]),
        ], styles))
        story.append(Spacer(1, 6))
    story.append(Spacer(1, 8))

    def _rep_section(label, snap):
        story.append(Paragraph(label, styles["Heading2"]))
        if not snap:
            story.append(Paragraph("Not recorded — test incomplete", styles["Normal"]))
            story.append(Spacer(1, 8))
            return
        story.append(_pdf_kv_table([
            ("Status", snap["status"]),
            ("Applied load", f"{_fmt(snap['applied_load'])} {snap['unit']}"),
            ("Readings", ", ".join(_fmt(x) for x in snap["readings"])),
            ("Weighing count", snap["weighing_count"]),
            ("Required weighings", snap["required_weighings"]),
            ("Minimum result", snap["minimum_result"]),
            ("Maximum result", snap["maximum_result"]),
            ("Spread", snap["spread"]),
            ("MPE", snap["mpe"]),
            ("Spread within MPE", snap["spread_within_mpe"]),
            ("Ruleset", snap["ruleset_version"]),
            ("Rule reference", snap["rule_reference"]),
        ], styles))
        story.append(Spacer(1, 8))

    story.append(Paragraph("Section 3.1: Eccentricity — Quarter Segments", styles["Heading2"]))
    ecc = data["eccentricity"]
    if not ecc:
        story.append(Paragraph("Not recorded — test incomplete", styles["Normal"]))
    else:
        pe = ecc.get("position_errors") or {}
        story.append(_pdf_kv_table([
            ("Status", ecc["status"]),
            ("Maximum additive tare effect", f"{_fmt(ecc['maximum_additive_tare_effect'])} {ecc['unit']}"),
            ("Required test load", f"{_fmt(ecc['required_test_load'])} {ecc['unit']}"),
            ("Required positions", ", ".join(ecc["required_positions"] or [])),
            ("Tested positions", ", ".join(ecc["tested_positions"] or [])),
            ("Q1 Ec", pe.get("Q1", "—")),
            ("Q2 Ec", pe.get("Q2", "—")),
            ("Q3 Ec", pe.get("Q3", "—")),
            ("Q4 Ec", pe.get("Q4", "—")),
            ("Failed positions", ", ".join(ecc["failed_positions"] or []) or "none"),
            ("MPE", ecc["mpe"]),
            ("Ruleset", ecc["ruleset_version"]),
            ("Rule reference", ecc["rule_reference"]),
        ], styles))
    story.append(Spacer(1, 10))
    _rep_section("Section 5: Repeatability — Half Max", data["repeatability"]["half_max"])
    _rep_section("Section 5: Repeatability — Near Max", data["repeatability"]["near_max"])
    story.append(Paragraph("Tare weighing", styles["Heading2"]))
    tw = data.get("tare_weighing")
    if not tw:
        story.append(Paragraph("Not recorded — test not performed", styles["Normal"]))
    else:
        story.append(_pdf_kv_table([
            ("Status", tw["status"]),
            ("Tare value", f"{_fmt(tw['tare_value'])} {tw['unit']}"),
            ("Observation count", tw["observation_count"]),
            ("Required observations", tw["required_observations"]),
            ("Complete", tw["complete"]),
            ("Failed observation indices", ", ".join(str(i) for i in (tw.get("failed_observation_indices") or [])) or "none"),
            ("Ruleset", tw["ruleset_version"]),
            ("Rule reference", tw["rule_reference"]),
        ], styles))
        headers_t = ["Net", "Gross", "L", "I", "ΔL", "P", "E", "E0", "Ec", "|Ec|", "MPE", "Margin", "Result", "Method"]
        table_t = [headers_t]
        for o in tw.get("observations") or []:
            table_t.append([
                _fmt(o.get("net_load")), _fmt(o.get("gross_load")), _fmt(o.get("load")), _fmt(o.get("indication")),
                _fmt(o.get("additional_load")), _fmt(o.get("indication_before_rounding")), _fmt(o.get("error_before_rounding")),
                _fmt(o.get("zero_error")), _fmt(o.get("corrected_error")), _fmt(o.get("absolute_corrected_error")),
                _fmt(o.get("mpe")), _fmt(o.get("margin_to_limit")), o.get("result") or "—", o.get("error_method") or "—",
            ])
        if len(table_t) == 1:
            table_t.append(["No observations"] + [""] * 13)
        tw_tbl = Table(table_t, colWidths=[32]*14, repeatRows=1)
        tw_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 6),
        ]))
        story.append(Spacer(1, 6))
        story.append(tw_tbl)
    story += [

        Spacer(1, 10),
        Paragraph("Standards and Rule Provenance", styles["Heading2"]),
        _pdf_kv_table([
            ("OIML R 76-1 edition", data["programme"]["standard_edition"]),
            ("Ruleset version", data["programme"]["ruleset_version"]),
            ("Collected rule references", "; ".join(data["programme"]["rule_references"]) or "None recorded"),
            ("Report format basis", data["format_basis"]),
            ("full_standard_coverage", "No"),
            ("Existing report status", data["status"]),
            ("Existing final_result field", data["final_result"]),
        ], styles),
    ]
    SimpleDocTemplate(
        path, pagesize=A4, rightMargin=28, leftMargin=28, topMargin=30, bottomMargin=28,
    ).build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)


def _docx_kv(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (a, b) in enumerate(rows):
        t.rows[i].cells[0].text = str(a)
        t.rows[i].cells[1].text = _fmt(b)


def build_docx(r, path):
    data = _build_report_document_data(r)
    d = Document()
    d.add_heading(data["title"], 0)
    d.add_paragraph(data["subtitle"])
    d.add_paragraph(data["programme"]["scope_note"])
    d.add_paragraph(data["scope_warning"])
    d.add_heading("General information concerning the type", 1)
    inst = data["instrument"]
    _docx_kv(d, [
        ("Report ID", data["report_id"]),
        ("Report number", data["report_number"]),
        ("Instrument ID", inst["id"]),
        ("Manufacturer", inst["manufacturer"]),
        ("Type designation / model", inst["model"]),
        ("Serial number", inst["serial_number"]),
        ("Accuracy class", inst["accuracy_class"]),
        ("Max", f"{_fmt(inst['max_capacity'])} {inst['unit']}"),
        ("Min", f"{_fmt(inst['min_capacity'])} {inst['unit']}"),
        ("Verification scale interval e", f"{_fmt(inst['e'])} {inst['unit']}"),
        ("Actual scale interval d", f"{_fmt(inst['d'])} {inst['unit']}"),
        ("Number of verification intervals n", _fmt(inst["n"])),
        ("Software version", inst["software_version"]),
        ("Report status", data["status"]),
        ("Created", data["created_at"]),
        ("Ruleset version", data["ruleset_version"]),
    ])
    d.add_heading("Summary of Implemented Test Evaluation", 1)
    d.add_paragraph("Overall PASS indicates only that all currently implemented core modules passed.")
    p = data["programme"]
    _docx_kv(d, [
        ("Overall implemented-core status", p["overall_status"]),
        ("Zero-setting", p.get("zero_setting_status") or "INCOMPLETE"),
        ("Observation / weighing checks", p["observation_status"]),
        ("Repeatability", p["repeatability_status"]),
        ("Eccentricity", p["eccentricity_status"]),
        ("Tare weighing", p.get("tare_weighing_status") or "INCOMPLETE"),
        ("Repeatability Half-Max", p["repeatability_half_max_status"] or "Not recorded"),
        ("Repeatability Near-Max", p["repeatability_near_max_status"] or "Not recorded"),
        ("Failed modules", ", ".join(p["failed_modules"]) or "none"),
        ("Incomplete modules", ", ".join(p["incomplete_modules"]) or "none"),
        ("Ruleset", p["ruleset_version"]),
        ("full_standard_coverage", "No"),
        ("Scope note", p["scope_note"]),
    ])
    d.add_heading("Zero-setting accuracy", 1)
    zs = data.get("zero_setting")
    if not zs:
        d.add_paragraph("Not recorded — test not performed")
    else:
        _docx_kv(d, [
            ("Status", zs["status"]),
            ("Setting mode", zs["setting_mode"]),
            ("Indication", f"{_fmt(zs['indication'])} {zs['unit']}"),
            ("Additional load", "—" if zs["additional_load"] is None else f"{_fmt(zs['additional_load'])} {zs['unit']}"),
            ("Error method", zs["error_method"]),
            ("Indication before rounding", f"{_fmt(zs['indication_before_rounding'])} {zs['unit']}"),
            ("Zero error", f"{_fmt(zs['zero_error'])} {zs['unit']}"),
            ("Absolute zero error", f"{_fmt(zs['absolute_zero_error'])} {zs['unit']}"),
            ("Permitted limit", f"{_fmt(zs['permitted_limit'])} {zs['unit']}"),
            ("Margin to limit", f"{_fmt(zs['margin_to_limit'])} {zs['unit']}"),
            ("e", f"{_fmt(zs['verification_scale_interval'])} {zs['unit']}"),
            ("d", f"{_fmt(zs['actual_scale_interval'])} {zs['unit']}"),
            ("Ruleset", zs["ruleset_version"]),
            ("Rule reference", zs["rule_reference"]),
        ])
    d.add_heading("Section 1: Weighing performance", 1)
    tbl = d.add_table(rows=1, cols=12)
    tbl.style = "Table Grid"
    heads = ["Test", "Pos", "Applied", "L", "I", "ΔL", "P", "E", "E0", "Ec", "mpe", "Result"]
    for cell, h in zip(tbl.rows[0].cells, heads):
        cell.text = h
    if not data["observations"]:
        row = tbl.add_row().cells
        row[0].text = "No observations recorded"
    for o in data["observations"]:
        cells = tbl.add_row().cells
        c = o["calculation"]
        if o["legacy"] or c is None:
            vals = [o["test_type"] or "—", o["position"] or "—", _fmt(o["applied_load"]),
                    _fmt(o["reference_value"]), _fmt(o["indication"]), "—", "—", "—", "—", "—", "—",
                    "Legacy / no OIML calculation snapshot"]
        else:
            vals = [o["test_type"] or "—", o["position"] or "—", _fmt(o["applied_load"]),
                    _fmt(c["canonical_load"]), _fmt(o["indication"]), _fmt(c["additional_load"]),
                    _fmt(c["indication_before_rounding"]), _fmt(c["error_before_rounding"]),
                    _fmt(c["zero_error"]), _fmt(c["corrected_error"]), _fmt(c["mpe"]), o["result"]]
        for cell, v in zip(cells, vals):
            cell.text = str(v)
        c = o["calculation"]
        if not o["legacy"] and c is not None:
            d.add_paragraph(
                f"|Ec|={_fmt(c['absolute_corrected_error'])}; margin={_fmt(c['margin_to_limit'])}; "
                f"method={c['error_method']}; ruleset={c['ruleset_version']}; "
                f"{c['error_rule_reference']}; {c['mpe_rule_reference']}"
            )

    def _rep(heading, snap):
        d.add_heading(heading, 1)
        if not snap:
            d.add_paragraph("Not recorded — test incomplete")
            return
        _docx_kv(d, [
            ("Status", snap["status"]),
            ("Applied load", f"{_fmt(snap['applied_load'])} {snap['unit']}"),
            ("Readings", ", ".join(_fmt(x) for x in snap["readings"])),
            ("Weighing count", snap["weighing_count"]),
            ("Required weighings", snap["required_weighings"]),
            ("Minimum result", snap["minimum_result"]),
            ("Maximum result", snap["maximum_result"]),
            ("Spread", snap["spread"]),
            ("MPE", snap["mpe"]),
            ("Spread within MPE", snap["spread_within_mpe"]),
            ("Ruleset", snap["ruleset_version"]),
            ("Rule reference", snap["rule_reference"]),
        ])

    d.add_heading("Section 3.1: Eccentricity — Quarter Segments", 1)
    ecc = data["eccentricity"]
    if not ecc:
        d.add_paragraph("Not recorded — test incomplete")
    else:
        pe = ecc.get("position_errors") or {}
        _docx_kv(d, [
            ("Status", ecc["status"]),
            ("Maximum additive tare effect", f"{_fmt(ecc['maximum_additive_tare_effect'])} {ecc['unit']}"),
            ("Required test load", f"{_fmt(ecc['required_test_load'])} {ecc['unit']}"),
            ("Required positions", ", ".join(ecc["required_positions"] or [])),
            ("Tested positions", ", ".join(ecc["tested_positions"] or [])),
            ("Q1 Ec", pe.get("Q1", "—")),
            ("Q2 Ec", pe.get("Q2", "—")),
            ("Q3 Ec", pe.get("Q3", "—")),
            ("Q4 Ec", pe.get("Q4", "—")),
            ("Failed positions", ", ".join(ecc["failed_positions"] or []) or "none"),
            ("MPE", ecc["mpe"]),
            ("Ruleset", ecc["ruleset_version"]),
            ("Rule reference", ecc["rule_reference"]),
        ])
    _rep("Section 5: Repeatability — Half Max", data["repeatability"]["half_max"])
    _rep("Section 5: Repeatability — Near Max", data["repeatability"]["near_max"])
    d.add_heading("Tare weighing", 1)
    tw = data.get("tare_weighing")
    if not tw:
        d.add_paragraph("Not recorded — test not performed")
    else:
        _docx_kv(d, [
            ("Status", tw["status"]),
            ("Tare value", f"{_fmt(tw['tare_value'])} {tw['unit']}"),
            ("Observation count", tw["observation_count"]),
            ("Required observations", tw["required_observations"]),
            ("Complete", tw["complete"]),
            ("Failed observation indices", ", ".join(str(i) for i in (tw.get("failed_observation_indices") or [])) or "none"),
            ("Ruleset", tw["ruleset_version"]),
            ("Rule reference", tw["rule_reference"]),
        ])
        tblt = d.add_table(rows=1, cols=8)
        tblt.style = "Table Grid"
        for cell, h in zip(tblt.rows[0].cells, ["Net", "Gross", "L", "I", "Ec", "MPE", "Result", "Method"]):
            cell.text = h
        for o in tw.get("observations") or []:
            cells = tblt.add_row().cells
            vals = [o.get("net_load"), o.get("gross_load"), o.get("load"), o.get("indication"),
                    o.get("corrected_error"), o.get("mpe"), o.get("result"), o.get("error_method")]
            for cell, v in zip(cells, vals):
                cell.text = _fmt(v)
            d.add_paragraph(
                f"ΔL={_fmt(o.get('additional_load'))}; P={_fmt(o.get('indication_before_rounding'))}; "
                f"E={_fmt(o.get('error_before_rounding'))}; E0={_fmt(o.get('zero_error'))}; "
                f"|Ec|={_fmt(o.get('absolute_corrected_error'))}; margin={_fmt(o.get('margin_to_limit'))}"
            )
    d.add_heading("Standards and Rule Provenance", 1)

    _docx_kv(d, [
        ("OIML R 76-1 edition", p["standard_edition"]),
        ("Ruleset version", p["ruleset_version"]),
        ("Collected rule references", "; ".join(p["rule_references"]) or "None recorded"),
        ("Report format basis", data["format_basis"]),
        ("full_standard_coverage", "No"),
        ("Existing report status", data["status"]),
        ("Existing final_result field", data["final_result"]),
    ])
    d.save(path)

@app.get("/api/reports/<int:rid>/pdf")
@require_role(*REPORT_READ_ROLES)
def pdf(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Not found"}),404
    denied=_assert_report_readable(r)
    if denied: return denied
    path=os.path.join(REPORTS,r.report_number+".pdf"); build_pdf(r,path); return send_file(path,as_attachment=True,download_name=os.path.basename(path),mimetype="application/pdf")

@app.get("/api/reports/<int:rid>/word")
@require_role(*REPORT_READ_ROLES)
def word(rid):
    r=db.session.get(Report,rid)
    if not r: return jsonify({"error":"Not found"}),404
    denied=_assert_report_readable(r)
    if denied: return denied
    path=os.path.join(REPORTS,r.report_number+".docx"); build_docx(r,path); return send_file(path,as_attachment=True,download_name=os.path.basename(path),mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

@app.get("/media/products/<path:name>")
def media_product(name): return send_from_directory(PRODUCT_IMAGES,name)
@app.get("/media/evidence/<path:name>")
@require_role(*REPORT_READ_ROLES)
def media_evidence(name):
    evidence=TestEvidence.query.filter_by(stored_name=name).first()
    if not evidence: return jsonify({"error":"Not found"}),404
    denied=_assert_report_readable(evidence.report)
    if denied: return denied
    return send_from_directory(EVIDENCE,name)
@app.get("/media/uploads/<path:name>")
@require_role(*REPORT_READ_ROLES)
def media_upload(name):
    attachment=Attachment.query.filter_by(stored_name=name).first()
    if not attachment: return jsonify({"error":"Not found"}),404
    denied=_assert_report_readable(attachment.report)
    if denied: return denied
    return send_from_directory(UPLOAD,name)


@app.cli.command("create-user")
@click.option("--username", prompt=True, help="Username for the new account.")
@click.option(
    "--role",
    type=click.Choice(("Admin", "Tester", "Reviewer", "Manufacturer"), case_sensitive=True),
    prompt=True,
    help="Role for the new account.",
)
@click.password_option(confirmation_prompt=True)
def create_user(username, role, password):
    username = username.strip()
    if not username:
        raise click.UsageError("Username must not be blank")
    if User.query.filter_by(username=username).first():
        raise click.ClickException(f"User '{username}' already exists")
    db.session.add(User(username=username, role=role, password_hash=generate_password_hash(password)))
    db.session.commit()
    click.echo(f"Created {role} user '{username}'.")

@app.route("/", defaults={"path":""})
@app.route("/<path:path>")
def spa(path):
    if path.startswith(("api/","media/")): return jsonify({"error":"Not found"}),404
    index=os.path.join(FRONTEND_DIST,"index.html")
    if os.path.exists(index): return send_from_directory(FRONTEND_DIST,"index.html")
    return "React frontend not built. Run: cd frontend && npm install && npm run build",503

with app.app_context():
    db.create_all()
    if app.config["DEMO_MODE"]:
        # Demo accounts are opt-in and are never reset or overwritten.
        seeds=[("admin","admin123","Admin"),("tester","tester123","Tester"),("reviewer","reviewer123","Reviewer"),("manufacturer","manufacturer123","Manufacturer")]
        for username,password,role in seeds:
            if not User.query.filter_by(username=username).first():
                db.session.add(User(username=username,password_hash=generate_password_hash(password),role=role))
        db.session.commit()

if __name__=="__main__":
    print("NAWI-TestPro API running at http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=False)
